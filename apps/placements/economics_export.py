"""«Экономика» файлом Excel: сводка, таблица по месяцам с графиками, размещения.

Те же числа, что на экране (`apps.placements.economics`): суммы — в евро с
копейками, месяцы и даты — датами Excel, чтобы по ним можно было считать.
Графики — родные графики Excel: деньги по месяцам столбиками с линией бюджета
и размещённые ссылки линией с числом у каждой точки, как на экране.
"""

import datetime as dt
import io
from decimal import Decimal
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.placements.economics import Report
from apps.placements.economics_view import STATUS_LABELS, State, issues

MONEY = "#,##0.00"
DATE = "DD.MM.YYYY"
MONTH = "MMM YYYY"
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
QUIET = Font(color="6E6E73")
OVER = Font(color="D03B3B", bold=True)
PERIOD_FILL = PatternFill("solid", start_color="E8F0FB", end_color="E8F0FB")
HEADER_FILL = PatternFill("solid", start_color="F2F2F4", end_color="F2F2F4")
# Цвета рядов — как на экране в светлой расцветке (seo/admin-palettes.css).
SPENT_COLOR, PLANNED_COLOR, PLACED_COLOR, BUDGET_COLOR = "2A78D6", "EB6834", "4A3AA7", "6E6E73"


def _eur(cents: int | None) -> Decimal | None:
    return None if cents is None else Decimal(cents) / 100


def _day(value: dt.datetime | None) -> dt.date | None:
    return timezone.localtime(value).date() if value is not None else None


def xlsx(state: State, report: Report) -> bytes:
    book = Workbook()
    # Пустой лист новой книги заменяем своими: у `active` в заглушках тип «может не быть».
    book.remove(book.worksheets[0])
    _summary(book.create_sheet("Сводка"), state, report)
    _months(book.create_sheet("По месяцам"), report)
    _placements(book.create_sheet("Размещения"), state, report)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _header(sheet: Any, row: int, titles: list[str]) -> None:
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _money(sheet: Any, row: int, column: int, cents: int | None) -> Any:
    cell = sheet.cell(row=row, column=column, value=_eur(cents))
    cell.number_format = MONEY
    return cell


def _summary(sheet: Any, state: State, report: Report) -> None:
    total = report.total
    sheet["A1"] = "Экономика размещений"
    sheet["A1"].font = TITLE
    generated = timezone.localtime()
    sheet["A2"] = f"{state.product_label} · {state.label} · отчёт на {generated:%d.%m.%Y %H:%M}"
    sheet["A2"].font = QUIET
    _header(sheet, 4, ["Показатель", "Значение", "Пояснение"])
    months = len(report.period.months)
    lines: list[tuple[str, int | None, str, bool]] = [
        (
            "Бюджет, €",
            total.budget_cents,
            f"{report.budget_month_cents / 100:,.0f} € в месяц × {months}".replace(",", " "),
            True,
        ),
        ("Потрачено, €", total.spent_cents, "«Заплачено» размещённых, по дате публикации", True),
        (
            "Запланировано, €",
            total.planned_cents,
            "в работе, заявка отправлена, написание статьи — по дате заявки",
            True,
        ),
        ("Осталось, €", total.left_cents, "бюджет − потрачено − запланировано", True),
        ("Размещено ссылок", total.placed, "опубликовано за период", False),
        ("Средняя цена, €", total.average_cents, "потрачено ÷ размещённые с суммой", True),
        ("В плане размещений", total.planned, "", False),
        ("Копирайт, €", total.copy_cents, "статьи наших копирайтеров — не из бюджета", True),
        ("Копирайт в работе, €", total.copy_planned_cents, "статьи, которые пишутся", True),
        ("Потрачено + копирайт, €", total.with_copy_cents, "во что обошлись размещения", True),
        (
            "Ср. цена + копирайт, €",
            total.average_with_copy_cents,
            "«потрачено + копирайт» ÷ размещённые с суммой",
            True,
        ),
    ]
    row = 5
    for title, value, note, money in lines:
        sheet.cell(row=row, column=1, value=title)
        if money:
            cell = _money(sheet, row, 2, value)
            if title.startswith("Осталось") and value is not None and value < 0:
                cell.font = OVER
        else:
            sheet.cell(row=row, column=2, value=value)
        sheet.cell(row=row, column=3, value=note).font = QUIET
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="План по статусам").font = BOLD
    row += 1
    _header(sheet, row, ["Статус", "Размещений", "Сумма, €"])
    for group in report.by_status:
        row += 1
        sheet.cell(row=row, column=1, value=group.name)
        sheet.cell(row=row, column=2, value=group.count)
        _money(sheet, row, 3, group.cents)

    row += 2
    sheet.cell(row=row, column=1, value="По продавцам").font = BOLD
    row += 1
    _header(sheet, row, ["Продавец", "Размещений", "Всего, €", "Потрачено, €", "Запланировано, €"])
    for group in report.by_seller:
        row += 1
        sheet.cell(row=row, column=1, value=group.name)
        sheet.cell(row=row, column=2, value=group.count)
        _money(sheet, row, 3, group.cents)
        _money(sheet, row, 4, group.spent_cents)
        _money(sheet, row, 5, group.planned_cents)

    if report.by_copywriter:
        row += 2
        sheet.cell(row=row, column=1, value="По копирайтерам").font = BOLD
        row += 1
        _header(sheet, row, ["Копирайтер", "Статей", "Всего, €", "Размещённые, €", "В работе, €"])
        for group in report.by_copywriter:
            row += 1
            sheet.cell(row=row, column=1, value=group.name)
            sheet.cell(row=row, column=2, value=group.count)
            _money(sheet, row, 3, group.cents)
            _money(sheet, row, 4, group.spent_cents)
            _money(sheet, row, 5, group.planned_cents)

    found = issues(state, report)
    if found:
        row += 2
        sheet.cell(row=row, column=1, value="Что не посчитано").font = BOLD
        for issue in found:
            row += 1
            sheet.cell(row=row, column=1, value=issue["count"])
            sheet.cell(row=row, column=2, value=issue["text"])
    for column, width in zip("ABCDE", (34, 16, 56, 16, 18), strict=True):
        sheet.column_dimensions[column].width = width


def _months(sheet: Any, report: Report) -> None:
    employees = [name or "без сотрудника" for name in report.employees]
    titles = [
        "Месяц",
        "Размещено ссылок",
        *employees,
        "Потрачено, €",
        "Ср. цена, €",
        "Копирайт, €",
        "Потрачено + копирайт, €",
        "Ср. цена + копирайт, €",
        "Запланировано, €",
        "Бюджет, €",
        "Осталось, €",
    ]
    _header(sheet, 1, titles)
    sheet.freeze_panes = "B2"
    first_money = 3 + len(employees)
    row = 1
    for month in report.months:
        row += 1
        in_period = month.month in report.period
        cell = sheet.cell(row=row, column=1, value=month.month)
        cell.number_format = MONTH
        sheet.cell(row=row, column=2, value=month.placed)
        for offset, name in enumerate(report.employees):
            sheet.cell(row=row, column=3 + offset, value=month.by_employee.get(name) or None)
        _money(sheet, row, first_money, month.spent_cents)
        _money(sheet, row, first_money + 1, month.average_cents)
        _money(sheet, row, first_money + 2, month.copy_cents)
        _money(sheet, row, first_money + 3, month.with_copy_cents)
        _money(sheet, row, first_money + 4, month.average_with_copy_cents)
        _money(sheet, row, first_money + 5, month.planned_cents)
        _money(sheet, row, first_money + 6, month.budget_cents)
        left = _money(sheet, row, first_money + 7, month.left_cents)
        if month.left_cents < 0:
            left.font = OVER
        if in_period:
            for column in range(1, len(titles) + 1):
                sheet.cell(row=row, column=column).fill = PERIOD_FILL
    last = row
    row += 1
    sheet.cell(row=row, column=1, value="Итого").font = BOLD
    sheet.cell(row=row, column=2, value=sum(m.placed for m in report.months)).font = BOLD
    for offset, name in enumerate(report.employees):
        value = sum(m.by_employee.get(name, 0) for m in report.months)
        sheet.cell(row=row, column=3 + offset, value=value).font = BOLD
    priced = sum(m.priced for m in report.months)
    spent = sum(m.spent_cents for m in report.months)
    planned = sum(m.planned_cents for m in report.months)
    budget = sum(m.budget_cents for m in report.months)
    copy = sum(m.copy_cents for m in report.months)
    for offset, cents in enumerate(
        (
            spent,
            round(spent / priced) if priced else None,
            copy,
            spent + copy,
            round((spent + copy) / priced) if priced else None,
            planned,
            budget,
            budget - spent - planned,
        )
    ):
        _money(sheet, row, first_money + offset, cents).font = BOLD
    sheet.column_dimensions["A"].width = 14
    for column in range(2, len(titles) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    sheet.row_dimensions[1].height = 32
    if last < 2:
        return
    anchor = get_column_letter(len(titles) + 2)
    months = Reference(sheet, min_col=1, min_row=2, max_row=last)

    links = LineChart()
    links.title = "Размещено ссылок по месяцам"
    links.height, links.width = 8, 22
    links.legend = None
    links.y_axis.title = None
    links.add_data(Reference(sheet, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    links.set_categories(months)
    links.x_axis.number_format = MONTH
    series = links.series[0]
    series.graphicalProperties.line.solidFill = PLACED_COLOR
    series.graphicalProperties.line.width = 22000
    series.marker.symbol = "circle"
    series.marker.size = 6
    series.marker.graphicalProperties.solidFill = PLACED_COLOR
    series.marker.graphicalProperties.line.solidFill = PLACED_COLOR
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.position = "t"
    sheet.add_chart(links, f"{anchor}2")

    money = BarChart()
    money.type = "col"
    money.grouping = "stacked"
    money.overlap = 100
    money.gapWidth = 60
    money.title = "Деньги по месяцам, €"
    money.height, money.width = 8, 22
    money.add_data(
        Reference(sheet, min_col=first_money, min_row=1, max_row=last), titles_from_data=True
    )
    money.add_data(
        Reference(sheet, min_col=first_money + 5, min_row=1, max_row=last), titles_from_data=True
    )
    money.set_categories(months)
    money.x_axis.number_format = MONTH
    money.series[0].graphicalProperties.solidFill = SPENT_COLOR
    money.series[0].graphicalProperties.line.noFill = True
    money.series[1].graphicalProperties.solidFill = PLANNED_COLOR
    money.series[1].graphicalProperties.line.noFill = True
    budget_line = LineChart()
    budget_line.add_data(
        Reference(sheet, min_col=first_money + 6, min_row=1, max_row=last), titles_from_data=True
    )
    budget_line.series[0].graphicalProperties.line.solidFill = BUDGET_COLOR
    budget_line.series[0].graphicalProperties.line.width = 19000
    budget_line.series[0].marker.symbol = "none"
    money += budget_line
    money.legend.position = "b"
    sheet.add_chart(money, f"{anchor}20")


def _placements(sheet: Any, state: State, report: Report) -> None:
    titles = [
        "Площадка",
        "Продукт",
        "Статус",
        "Продавец",
        "Сотрудник",
        "Заявка отправлена",
        "Опубликовано",
        "Месяц в отчёте",
        "Сумма, €",
        "Откуда сумма",
        "Заплачено в валюте",
        "Валюта",
        "Копирайтер",
        "Копирайт, €",
    ]
    _header(sheet, 1, titles)
    sheet.freeze_panes = "B2"
    row = 1
    for item in report.rows:
        row += 1
        sheet.cell(row=row, column=1, value=item.domain)
        sheet.cell(row=row, column=2, value=item.product)
        sheet.cell(row=row, column=3, value=STATUS_LABELS.get(item.status, item.status))
        sheet.cell(row=row, column=4, value=item.seller or None)
        sheet.cell(row=row, column=5, value=item.employee or None)
        for column, value in ((6, _day(item.ordered_at)), (7, _day(item.published_at))):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.number_format = DATE
        month = sheet.cell(row=row, column=8, value=item.month)
        month.number_format = MONTH
        _money(sheet, row, 9, item.eur_cents)
        if item.eur_cents is None:
            source = "нет суммы"
        elif item.from_paid:
            source = "заплачено" if item.spent else "договорено"
        else:
            source = "по рабочей цене площадки"
        sheet.cell(row=row, column=10, value=source)
        _money(sheet, row, 11, item.paid_cents)
        sheet.cell(row=row, column=12, value=item.currency if item.paid_cents is not None else None)
        sheet.cell(row=row, column=13, value=item.copywriter or None)
        _money(sheet, row, 14, item.copy_cents)
    widths = (30, 12, 20, 24, 14, 12, 12, 12, 12, 24, 14, 8, 18, 12)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    if row > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{row}"
