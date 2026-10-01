"""Чтение загруженного файла: таблица строк, строка заголовков, колонки.

Прайсы приходят в xlsx и csv: CSV — в UTF-8 (с BOM и без), Windows-1251
или UTF-16 (так выгружает Ahrefs), с разделителем `,`, `;` или табуляцией.
Над заголовком бывают заголовок листа и пустые строки — строка заголовков
ищется сама. Колонка без заголовка, но со значениями получает имя по букве
(«Колонка E»): из файла ничего не теряется (ADR-041).
"""

import csv
import datetime as dt
import io
import re
import warnings
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Строку заголовков ищем среди первых строк файла.
HEADER_SEARCH_ROWS = 20
# Сколько первых значений колонки показать человеку на шаге разметки.
SAMPLES = 3

_SPACES = re.compile(r"\s+")


class FileError(Exception):
    """Файл не прочитать: формат, кодировка, пустой. Текст — для человека."""


@dataclass(frozen=True)
class Column:
    """Колонка файла. `key` — по нему разметка запоминается у продавца."""

    index: int
    letter: str
    header: str  # как в файле; у колонки без заголовка — «Колонка E»
    key: str
    samples: tuple[str, ...]
    filled: int  # непустых значений в данных

    @property
    def is_empty(self) -> bool:
        return self.filled == 0


@dataclass(frozen=True)
class Row:
    line: int  # номер строки в файле — по нему человек найдёт её
    cells: tuple[object, ...]

    def get(self, index: int) -> object:
        return self.cells[index] if index < len(self.cells) else None


@dataclass(frozen=True)
class Table:
    header_row: int
    columns: tuple[Column, ...]
    rows: tuple[Row, ...]  # только непустые строки под заголовком
    blank_rows: int


def header_key(header: str) -> str:
    """Заголовок для сравнения: без регистра и лишних пробелов."""
    return _SPACES.sub(" ", header.replace("\xa0", " ")).strip().lower()


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def show(value: object) -> str:
    """Значение ячейки строкой: как его увидит человек и как оно ляжет в «прочие данные»."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return f"{value:%d.%m.%Y}" if value.time() == dt.time.min else f"{value:%d.%m.%Y %H:%M}"
    if isinstance(value, dt.date):
        return f"{value:%d.%m.%Y}"
    return str(value).strip()


def read_rows(path: Path) -> list[tuple[object, ...]]:
    """Все строки первого непустого листа xlsx или файла csv, как есть."""
    if not path.is_file():
        raise FileError(f"Файл не найден: {path.name}")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if suffix in (".csv", ".txt", ".tsv"):
        return _read_csv(path.read_bytes())
    raise FileError(f"Формат «{suffix or path.name}» не поддерживается — нужен xlsx или csv.")


def read_table(
    path: Path,
    *,
    header_row: int | None = None,
    is_header: Callable[[Sequence[str]], bool] | None = None,
) -> Table:
    """Таблица файла. Строка заголовков — заданная или найденная сама.

    `is_header` — узнаёт заголовок по значениям ячеек (например, есть
    колонка «площадка» по словарю синонимов). Не узнал ни одну строку —
    берётся первая строка хотя бы с двумя непустыми ячейками.
    """
    rows = read_rows(path)
    if not any(not is_blank(cell) for row in rows for cell in row):
        raise FileError("Файл пустой.")
    if header_row is None:
        header_row = find_header_row(rows, is_header)
    elif not 1 <= header_row <= len(rows):
        raise FileError(f"В файле нет строки {header_row}.")
    header_cells = rows[header_row - 1]
    data: list[Row] = []
    blank = 0
    for number, cells in enumerate(rows[header_row:], start=header_row + 1):
        if all(is_blank(cell) for cell in cells):
            blank += 1
            continue
        data.append(Row(number, tuple(cells)))
    width = max([len(header_cells), *(_used_width(row.cells) for row in data)])
    columns = _columns(header_cells, data, width)
    return Table(header_row, columns, tuple(data), blank)


def find_header_row(
    rows: Sequence[Sequence[object]], is_header: Callable[[Sequence[str]], bool] | None
) -> int:
    candidates: list[int] = []
    for number, cells in enumerate(rows[:HEADER_SEARCH_ROWS], start=1):
        texts = [show(cell) for cell in cells if not is_blank(cell)]
        # Строка с одной ячейкой — заголовок листа («sites»), а не колонок.
        if len(texts) < 2:
            continue
        if is_header is not None and is_header(texts):
            return number
        candidates.append(number)
    if not candidates:
        raise FileError(
            "Не нашлась строка заголовков: в первых строках нет двух заполненных ячеек."
        )
    return candidates[0]


def _used_width(cells: Sequence[object]) -> int:
    for index in range(len(cells), 0, -1):
        if not is_blank(cells[index - 1]):
            return index
    return 0


def _columns(header_cells: Sequence[object], data: Sequence[Row], width: int) -> tuple[Column, ...]:
    columns: list[Column] = []
    seen: dict[str, int] = {}
    for index in range(width):
        letter = get_column_letter(index + 1)
        values = [row.get(index) for row in data]
        filled = [show(value) for value in values if not is_blank(value)]
        header = show(header_cells[index]) if index < len(header_cells) else ""
        if not header:
            if not filled:
                # Ни заголовка, ни значений — пустой разделитель, его нет.
                continue
            header = f"Колонка {letter}"
        key = header_key(header)
        # Два одинаковых заголовка — разные колонки: «Price» и «Price (2)».
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            header = f"{header} ({seen[key]})"
            key = header_key(header)
        columns.append(
            Column(
                index=index,
                letter=letter,
                header=header,
                key=key,
                samples=tuple(filled[:SAMPLES]),
                filled=len(filled),
            )
        )
    return tuple(columns)


def _read_xlsx(path: Path) -> list[tuple[object, ...]]:
    try:
        with warnings.catch_warnings():
            # Проверку данных и условное форматирование openpyxl не читает и
            # предупреждает об этом — нам они не нужны.
            warnings.filterwarnings("ignore", category=UserWarning)
            workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise FileError(f"Не удалось открыть xlsx: {error}") from error
    try:
        for sheet in workbook.worksheets:
            rows: list[tuple[object, ...]] = [
                tuple(row) for row in sheet.iter_rows(values_only=True)
            ]
            if any(not is_blank(cell) for row in rows for cell in row):
                return rows
        return []
    finally:
        workbook.close()


def _read_csv(raw: bytes) -> list[tuple[object, ...]]:
    text = _decode(raw)
    sample = "\n".join(text.splitlines()[:HEADER_SEARCH_ROWS])
    delimiter = _delimiter(sample)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    return [tuple(row) for row in reader]


def _decode(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        # UTF-16 с меткой порядка байт — выгрузки Ahrefs.
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        # Не UTF-8 — значит, Excel под Windows сохранил в Windows-1251.
        return raw.decode("cp1251", errors="replace")


def _delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        # Sniffer не уверен (одна колонка, кавычки) — разделитель, которого больше.
        counts = {d: sample.count(d) for d in (";", "\t", ",")}
        return max(counts, key=lambda d: counts[d]) if any(counts.values()) else ","
