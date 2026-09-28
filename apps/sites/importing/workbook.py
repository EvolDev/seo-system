"""Чтение книги Excel: вкладка → строки «заголовок → значение».

Колонки ищутся по заголовку, а не по позиции: в таблицу добавляют
колонки («Новая?»), и позиции съезжают. Нет обязательной колонки —
`ImportAbort`, импорт не начинается.
"""

import datetime as dt
import re
import warnings
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from apps.sites.importing.values import is_blank, text

# «Pos 22.09.26» — позиции ключа на дату (маппинг §2.2).
_POSITION_COLUMN = re.compile(r"Pos (\d{2})\.(\d{2})\.(\d{2})")


class ImportAbort(Exception):
    """Импорт нельзя продолжать. Ничего не записано: транзакция откатится."""


@dataclass(frozen=True)
class SheetRow:
    number: int  # номер строки в Excel — по нему человек найдёт её в таблице
    values: dict[str, object]

    def get(self, column: str) -> object:
        return self.values.get(column)


@dataclass(frozen=True)
class Sheet:
    name: str
    headers: tuple[str, ...]
    rows: tuple[SheetRow, ...]
    blank_rows: int


@contextmanager
def open_workbook(path: Path) -> Iterator[Workbook]:
    """Книга только для чтения; формулы — их сохранённые значения.

    `@contextmanager` превращает генератор в объект для `with`: код до
    `yield` — вход в блок, после — выход, даже если внутри было
    исключение. Так файл закрывается всегда.
    """
    if not path.is_file():
        raise ImportAbort(f"Файл не найден: {path}")
    with warnings.catch_warnings():
        # Проверку данных в ячейках (выпадающие списки таблицы) openpyxl не
        # читает и предупреждает об этом на каждой вкладке. Нам она не нужна.
        warnings.filterwarnings("ignore", "Unknown extension", UserWarning)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            yield workbook
        finally:
            workbook.close()


def read_sheet(workbook: Workbook, name: str, required: Iterable[str]) -> Sheet:
    """Строки вкладки. Полностью пустые строки не возвращаются, только считаются."""
    if name not in workbook.sheetnames:
        raise ImportAbort(f"В файле нет вкладки «{name}»")
    rows = workbook[name].iter_rows(values_only=True)
    header_cells = next(rows, ())
    headers = tuple(text(cell) or "" for cell in header_cells)
    missing = [column for column in required if column not in headers]
    if missing:
        raise ImportAbort(f"Во вкладке «{name}» нет колонок: {', '.join(missing)}")
    repeated = sorted({h for h in headers if h and headers.count(h) > 1})
    if repeated:
        raise ImportAbort(f"Во вкладке «{name}» повторяются колонки: {', '.join(repeated)}")

    result: list[SheetRow] = []
    blank = 0
    # Строка 1 — заголовок, данные начинаются со второй.
    for number, cells in enumerate(rows, start=2):
        values: dict[str, object] = {
            header: cell for header, cell in zip(headers, cells, strict=False) if header
        }
        if all(is_blank(value) for value in values.values()):
            blank += 1
            continue
        result.append(SheetRow(number, values))
    return Sheet(name, headers, tuple(result), blank)


def position_columns(headers: Iterable[str]) -> dict[str, dt.date]:
    """Колонки позиций по шаблону имени: «Pos 22.09.26» → 22.09.2026."""
    found: dict[str, dt.date] = {}
    for header in headers:
        match = _POSITION_COLUMN.fullmatch(header)
        if match:
            day, month, year = (int(part) for part in match.groups())
            found[header] = dt.date(2000 + year, month, day)
    return found
