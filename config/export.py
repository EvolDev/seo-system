"""Выгрузка списка в файл для человека: Excel и CSV (E1-06, ADR-053).

Экран описывает колонки (`Column`: заголовок, вид значения, ширина) и отдаёт
строки значений как есть: деньги — целые центы, даты — моменты или дни,
пусто — None. Как это выглядит в файле, решает этот модуль:

- Excel — настоящий .xlsx: первая строка закреплена и с автофильтром, деньги —
  числом с копейками, даты — датой «дд.мм.гггг», да/нет — «Да» и «Нет», как в
  таблице линкбилдинга;
- CSV — для Excel с русскими настройками: UTF-8 с BOM (иначе кириллица
  кракозябрами), `;` между колонками, запятая в дробях. Только первый лист.

Текст из чужих файлов (анкоры, комментарии) никогда не становится формулой:
в Excel ячейка — строка, в CSV перед «=», «+», «-», «@» стоит апостроф.

Кнопка «Выгрузить» над списком админки открывает окно выгрузки
(`admin/seo_export_tools.html`, тег `export_tools` — `config/export_tags.py`,
скрипт `seo/export.js`): какие колонки выгружать — по умолчанию все, снятые
галочки запоминаются в сессии у каждого списка своими; Excel или CSV. Строки —
тот же отбор и порядок, что у списка (фильтры, поиск, сортировка), все
страницы; отмечены строки галочками — только они.
"""

import csv
import datetime as dt
import io
from collections.abc import Callable, Iterable, Iterator, Sequence, Set
from dataclasses import dataclass
from decimal import Decimal
from enum import StrEnum
from typing import Any

from django.contrib.admin import ModelAdmin
from django.contrib.admin.options import IncorrectLookupParameters
from django.contrib.admin.views.main import ChangeList
from django.core.exceptions import PermissionDenied
from django.db.models import QuerySet
from django.http import (
    Http404,
    HttpRequest,
    HttpResponse,
    HttpResponseBadRequest,
    HttpResponseRedirect,
)
from django.urls import reverse
from django.utils import timezone
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Value = str | int | Decimal | bool | dt.date | Sequence[str] | None


class Kind(StrEnum):
    TEXT = "text"
    INT = "int"
    MONEY = "money"  # целые центы → сумма с копейками
    NUMBER = "number"  # Decimal как есть: доля, сумма в евро после пересчёта
    DATE = "date"  # день; момент — по времени проекта
    MONTH = "month"  # месяц дня: «сентябрь 2026», в Excel — дата первого числа


class Format(StrEnum):
    XLSX = "xlsx"
    CSV = "csv"


@dataclass(frozen=True)
class Column:
    title: str
    kind: Kind = Kind.TEXT
    width: int = 14  # ширина в Excel, в знаках


@dataclass(frozen=True)
class Sheet:
    title: str
    columns: Sequence[Column]
    rows: Iterable[Sequence[Value]]


@dataclass(frozen=True)
class Choice:
    """Галочка в окне выгрузки: одна колонка файла или группа (анкоры и ссылки)."""

    key: str
    title: str
    group: str  # заголовок раздела в окне: «Метрики», «Рабочая цена»…


# Раскладка листа: колонка файла и ключ галочки, которая её включает.
Layout = Sequence[tuple[str, Column]]


def narrow(
    layout: Layout, rows: Iterable[Sequence[Value]], wanted: Set[str]
) -> tuple[list[Column], Iterator[list[Value]]]:
    """Только выбранные колонки — в порядке раскладки, а не щелчков."""
    keep = [index for index, (key, _) in enumerate(layout) if key in wanted]
    columns = [layout[index][1] for index in keep]
    return columns, ([row[index] for index in keep] for row in rows)


CONTENT_TYPES = {
    Format.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    Format.CSV: "text/csv; charset=utf-8",
}

_MONEY_FORMAT = "#,##0.00"
_NUMBER_FORMAT = "#,##0.00"
_DATE_FORMAT = "DD.MM.YYYY"
_MONTH_FORMAT = "MMMM YYYY"
_HEADER_FONT = Font(bold=True)
# Так начинается формула в Excel; OWASP «CSV Injection».
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_HUNDRED = Decimal(100)


class ExportChangeList(ChangeList):
    """Тот же отбор и порядок, что у списка, но без страниц: выгрузке нужна вся
    выборка (`queryset`), а не строки одной страницы и их подсчёт."""

    def get_results(self, request: HttpRequest) -> None:
        self.result_list: list[Any] = []
        self.result_count = 0
        self.full_result_count = None
        self.show_full_result_count = False
        self.show_admin_actions = False
        self.can_show_all = False
        self.multi_page = False


def list_key(model_admin: "ModelAdmin[Any]") -> str:
    """Список, чей выбор колонок помнит сессия: `sites_productsitelatest`."""
    opts = model_admin.model._meta
    return f"{opts.app_label}_{opts.model_name}"


def url_name(model_admin: "ModelAdmin[Any]") -> str:
    """Имя адреса выгрузки списка: `sites_productsitelatest_export`."""
    return f"{list_key(model_admin)}_export"


def is_export(request: HttpRequest) -> bool:
    """Запрос — выгрузка, а не список: `get_changelist` отдаёт `ExportChangeList`."""
    match = request.resolver_match
    return match is not None and (match.url_name or "").endswith("_export")


# Снятые галочки по спискам: {"sites_productsitelatest": ["dr", …]}. Храним снятые,
# а не отмеченные: колонка, которая появится позже, сразу будет выгружаться.
SESSION_KEY = "seo_export_off"

# Листы и имя файла (без расширения) по списку, его выборке и выбранным ключам.
Build = Callable[[ChangeList, QuerySet[Any], Set[str]], tuple[Sequence[Sheet], str]]


def unchecked(request: HttpRequest, key: str) -> set[str]:
    """Колонки списка, с которых в прошлый раз сняли галочку."""
    stored = request.session.get(SESSION_KEY)
    values = stored.get(key, []) if isinstance(stored, dict) else []
    return {str(value) for value in values}


def _remember(request: HttpRequest, key: str, keys: Set[str], wanted: Set[str]) -> None:
    # Галочки, которых сейчас нет в окне (колонки региона без региона), — как были.
    off = (unchecked(request, key) - keys) | (keys - wanted)
    stored = request.session.get(SESSION_KEY)
    changed = dict(stored) if isinstance(stored, dict) else {}
    changed[key] = sorted(off)
    request.session[SESSION_KEY] = changed


def export_view(
    model_admin: "ModelAdmin[Any]",
    request: HttpRequest,
    fmt: str,
    choices: Sequence[Choice],
    build: Build,
) -> HttpResponse:
    """Файл по окну выгрузки (POST) или по ссылке без скрипта (GET).

    POST: `columns` — отмеченные галочки, `ids` — отмеченные строки списка
    (нет или «выбраны все N» — все отобранные). GET — все строки и колонки,
    как запомнила сессия.
    """
    if fmt not in Format:
        raise Http404
    if not model_admin.has_view_permission(request):
        raise PermissionDenied
    key = list_key(model_admin)
    keys = {choice.key for choice in choices}
    picked: list[int] = []
    if request.method == "POST":
        wanted = set(request.POST.getlist("columns")) & keys
        if not wanted:
            return HttpResponseBadRequest("Отметьте хотя бы одну колонку.")
        _remember(request, key, keys, wanted)
        if request.POST.get("select_across") != "1":
            picked = [int(value) for value in request.POST.getlist("ids") if value.isdigit()]
    else:
        wanted = (keys - unchecked(request, key)) or keys
    try:
        changelist = model_admin.get_changelist_instance(request)
    except IncorrectLookupParameters:
        # Как у списка: непонятный фильтр в адресе — назад к списку.
        opts = model_admin.model._meta
        return HttpResponseRedirect(reverse(f"admin:{opts.app_label}_{opts.model_name}_changelist"))
    queryset: QuerySet[Any] = changelist.queryset
    if picked:
        # Отмеченные строки — внутри того же отбора и в том же порядке.
        queryset = queryset.filter(pk__in=picked)
    sheets, name = build(changelist, queryset, wanted)
    if picked:
        name = f"{name} — выбрано {len(picked)}"
    return response(sheets, name, Format(fmt))


def response(sheets: Sequence[Sheet], name: str, fmt: Format) -> HttpResponse:
    """Файл-ответ: имя без расширения, расширение — по формату."""
    content = xlsx(sheets) if fmt is Format.XLSX else csv_bytes(sheets[0])
    return attachment(content, f"{name}.{fmt.value}", CONTENT_TYPES[fmt])


def attachment(content: bytes, file_name: str, content_type: str) -> HttpResponse:
    response = HttpResponse(content, content_type=content_type)
    # Имя файла бывает с кириллицей — заголовок по RFC 6266 собирает Django.
    response["Content-Disposition"] = content_disposition_header(True, file_name)
    return response


def xlsx(sheets: Sequence[Sheet]) -> bytes:
    # write_only — строки уходят в файл по одной, без модели всей книги в памяти.
    book = Workbook(write_only=True)
    for sheet in sheets:
        _write_sheet(book.create_sheet(sheet.title[:31]), sheet)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# Лист потоковой записи (`WriteOnlyWorksheet`): ширины колонок и автофильтр у него
# есть при запуске, но нет в заглушках типов (types-openpyxl) — отсюда Any.
def _write_sheet(worksheet: Any, sheet: Sheet) -> None:
    for number, column in enumerate(sheet.columns, start=1):
        worksheet.column_dimensions[get_column_letter(number)].width = column.width
    worksheet.freeze_panes = "A2"
    header = []
    for column in sheet.columns:
        cell = WriteOnlyCell(worksheet, value=column.title)
        cell.font = _HEADER_FONT
        header.append(cell)
    worksheet.append(header)
    count = 0
    for row in sheet.rows:
        worksheet.append(
            [
                _xlsx_cell(worksheet, column, value)
                for column, value in zip(sheet.columns, row, strict=True)
            ]
        )
        count += 1
    last = get_column_letter(max(len(sheet.columns), 1))
    worksheet.auto_filter.ref = f"A1:{last}{count + 1}"


def _xlsx_cell(worksheet: Any, column: Column, value: Value) -> object:
    if value is None:
        return None
    if column.kind is Kind.MONEY and isinstance(value, int) and not isinstance(value, bool):
        return _formatted(worksheet, Decimal(value) / _HUNDRED, _MONEY_FORMAT)
    if column.kind is Kind.NUMBER and isinstance(value, Decimal | int):
        return _formatted(worksheet, value, _NUMBER_FORMAT)
    if column.kind is Kind.DATE and isinstance(value, dt.date):
        return _formatted(worksheet, _day(value), _DATE_FORMAT)
    if column.kind is Kind.MONTH and isinstance(value, dt.date):
        return _formatted(worksheet, _day(value).replace(day=1), _MONTH_FORMAT)
    if column.kind is Kind.INT and isinstance(value, int) and not isinstance(value, bool):
        return value
    text = _text(value)
    if not text:
        return None
    if text.startswith("="):
        # Иначе openpyxl запишет формулу: строка «=…» из чужого файла.
        cell = WriteOnlyCell(worksheet, value=text)
        cell.data_type = "s"
        return cell
    return text


def _formatted(worksheet: Any, value: Decimal | int | dt.date, number_format: str) -> object:
    # Decimal и date openpyxl пишет числом и датой, хотя в заглушках их нет.
    cell = WriteOnlyCell(worksheet, value=value)  # type: ignore[arg-type]
    cell.number_format = number_format
    return cell


def csv_bytes(sheet: Sheet) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([column.title for column in sheet.columns])
    for row in sheet.rows:
        writer.writerow(
            [_csv_value(column, value) for column, value in zip(sheet.columns, row, strict=True)]
        )
    # BOM в начале — по нему Excel узнаёт UTF-8.
    return buffer.getvalue().encode("utf-8-sig")


def _csv_value(column: Column, value: Value) -> str:
    if value is None:
        return ""
    if column.kind is Kind.MONEY and isinstance(value, int) and not isinstance(value, bool):
        return _decimal_comma(Decimal(value) / _HUNDRED)
    if column.kind is Kind.NUMBER and isinstance(value, Decimal | int):
        return _decimal_comma(Decimal(value))
    if column.kind is Kind.DATE and isinstance(value, dt.date):
        return f"{_day(value):%d.%m.%Y}"
    if column.kind is Kind.MONTH and isinstance(value, dt.date):
        return month_name(_day(value))
    if column.kind is Kind.INT and isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    text = _text(value)
    return f"'{text}" if text.startswith(_FORMULA_START) else text


def _decimal_comma(value: Decimal) -> str:
    # Без разделителя тысяч: «1 234,56» Excel прочитал бы текстом.
    return f"{value:.2f}".replace(".", ",")


def _day(value: dt.date) -> dt.date:
    if isinstance(value, dt.datetime):
        return timezone.localtime(value).date() if timezone.is_aware(value) else value.date()
    return value


MONTHS = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)


def month_name(day: dt.date) -> str:
    """«сентябрь 2026»."""
    return f"{MONTHS[day.month - 1]} {day.year}"


def _text(value: Value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, str):
        return value
    if isinstance(value, dt.date):
        return f"{_day(value):%d.%m.%Y}"
    if isinstance(value, int | Decimal):
        return str(value)
    return ", ".join(str(item) for item in value)
