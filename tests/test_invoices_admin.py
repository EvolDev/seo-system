"""Экран «Счета» и счёт из «Размещений» (E1-14, ADR-055)."""

import datetime as dt
import io
from typing import Any

import pytest
from django.contrib.auth.models import User
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.placements.models import Invoice, InvoiceItem, InvoiceStatus, Placement
from apps.sites.models import Product, Seller, Site

pytestmark = pytest.mark.django_db

PARTIAL = {"X-Seo-Partial": "1"}
ADD_URL = reverse("admin:placements_invoice_add")
LIST_URL = reverse("admin:placements_invoice_changelist")
PLACEMENTS_URL = reverse("admin:placements_placement_changelist")


@pytest.fixture
def product() -> Product:
    return Product.objects.create(name="Convertio", domain="convertio.co")


@pytest.fixture
def seller() -> Seller:
    return Seller.objects.create(name="StarMedia")


def _placement(product: Product, domain: str, **fields: Any) -> Placement:
    return Placement.objects.create(
        site=Site.objects.create(domain=domain), product=product, **fields
    )


def _form(
    seller: Seller,
    rows: list[tuple[Placement, str]],
    *,
    amount: str = "100",
    initial: list[InvoiceItem] | None = None,
    **fields: str,
) -> dict[str, str]:
    """Форма счёта, как её отправляет браузер: строки — размещение и доля."""
    initial = initial or []
    data = {
        "seller": str(seller.pk),
        "status": "issued",
        "amount_cents": amount,
        "currency": "EUR",
        "pay_url": "",
        "number": "",
        "issued_on": "2026-10-04",
        "paid_on": "",
        "paid_by": "",
        "comment": "",
        "items-TOTAL_FORMS": str(len(rows)),
        "items-INITIAL_FORMS": str(len(initial)),
        "items-MIN_NUM_FORMS": "0",
        "items-MAX_NUM_FORMS": "1000",
    }
    for index, (placement, share) in enumerate(rows):
        data[f"items-{index}-placement"] = str(placement.pk)
        data[f"items-{index}-amount_cents"] = share
        if index < len(initial):
            data[f"items-{index}-id"] = str(initial[index].pk)
            data[f"items-{index}-invoice"] = str(initial[index].invoice_id)
    data.update(fields)
    return data


def _paid(placement: Placement) -> tuple[int | None, str | None]:
    placement.refresh_from_db()
    return placement.price_paid_cents, placement.currency


class TestAddInvoice:
    def test_batch_split_evenly_and_paid_set(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        rows = [_placement(product, f"site{n}.com") for n in range(3)]
        data = _form(seller, [(p, "") for p in rows])
        response = admin_client.post(ADD_URL, data, headers=PARTIAL)
        assert response.json()["saved"] is True
        invoice = Invoice.objects.get()
        assert sorted(invoice.items.values_list("amount_cents", flat=True)) == [3333, 3333, 3334]
        assert sorted(_paid(p)[0] or 0 for p in rows) == [3333, 3333, 3334]
        # Продавца размещению без продавца ставит счёт.
        assert {p.seller_id for p in Placement.objects.all()} == {seller.pk}

    def test_shares_short_not_saved(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        first, second = _placement(product, "a.com"), _placement(product, "b.com")
        response = admin_client.post(ADD_URL, _form(seller, [(first, "60"), (second, "38")]))
        assert response.status_code == 200
        assert "Доли меньше суммы счёта на €2" in response.content.decode()
        assert not Invoice.objects.exists()
        assert _paid(first)[0] is None

    def test_other_seller_not_saved(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        other = Seller.objects.create(name="Authlinker")
        placement = _placement(product, "dev.to", seller=other)
        response = admin_client.post(ADD_URL, _form(seller, [(placement, "")]))
        assert "куплено у продавца «Authlinker»" in response.content.decode()
        assert not Invoice.objects.exists()

    def test_same_placement_twice(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        placement = _placement(product, "dev.to")
        response = admin_client.post(ADD_URL, _form(seller, [(placement, ""), (placement, "")]))
        assert "Размещение в счёте дважды" in response.content.decode()

    def test_paid_status_fills_date_and_payer(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        placement = _placement(product, "dev.to")
        admin_client.post(ADD_URL, _form(seller, [(placement, "")], status="paid"))
        invoice = Invoice.objects.get()
        assert invoice.paid_on == timezone.localdate()
        assert invoice.paid_by == User.objects.get(username="admin")


class TestEditInvoice:
    @pytest.fixture
    def batch(self, admin_client: Client, product: Product, seller: Seller) -> Invoice:
        rows = [_placement(product, f"site{n}.com") for n in range(2)]
        admin_client.post(ADD_URL, _form(seller, [(p, "") for p in rows]))
        return Invoice.objects.get()

    def _url(self, invoice: Invoice) -> str:
        return reverse("admin:placements_invoice_change", args=[invoice.pk])

    def _rows(self, invoice: Invoice) -> list[InvoiceItem]:
        return list(invoice.items.order_by("pk"))

    def test_remove_placement_clears_its_paid(
        self, admin_client: Client, seller: Seller, batch: Invoice
    ) -> None:
        items = self._rows(batch)
        data = _form(
            seller,
            [(item.placement, "") for item in items],
            initial=items,
            **{"items-1-DELETE": "on"},
        )
        response = admin_client.post(self._url(batch), data, headers=PARTIAL)
        assert response.json()["saved"] is True
        assert list(batch.items.values_list("amount_cents", flat=True)) == [10000]
        assert _paid(items[0].placement)[0] == 10000
        assert _paid(items[1].placement)[0] is None

    def test_cancel_clears_paid(self, admin_client: Client, seller: Seller, batch: Invoice) -> None:
        items = self._rows(batch)
        rows = [(item.placement, f"{item.amount_cents / 100:.2f}") for item in items]
        data = _form(seller, rows, initial=items, status="cancelled")
        admin_client.post(self._url(batch), data)
        batch.refresh_from_db()
        assert batch.status == InvoiceStatus.CANCELLED
        assert [_paid(item.placement)[0] for item in items] == [None, None]
        # Строки остаются — история: что закрывал отменённый счёт.
        assert batch.items.count() == 2

    def test_new_total_with_old_shares_is_error(
        self, admin_client: Client, seller: Seller, batch: Invoice
    ) -> None:
        items = self._rows(batch)
        rows = [(item.placement, "50,00") for item in items]
        data = _form(seller, rows, initial=items, amount="120")
        response = admin_client.post(self._url(batch), data)
        assert "Доли меньше суммы счёта на €20" in response.content.decode()

    def test_paid_invoice_rows_locked(self, admin_client: Client, batch: Invoice) -> None:
        batch.status = InvoiceStatus.PAID
        batch.paid_on = timezone.localdate()
        batch.save()
        page = admin_client.get(self._url(batch), headers=PARTIAL).content.decode()
        assert 'name="items-0-amount_cents"' not in page
        assert 'name="items-0-DELETE"' not in page
        assert 'name="amount_cents"' not in page
        assert "Поделить поровну" not in page
        assert "€100" in page


class TestListAndAction:
    def test_mark_paid(self, admin_client: Client, product: Product, seller: Seller) -> None:
        issued = Invoice.objects.create(seller=seller, amount_cents=100)
        cancelled = Invoice.objects.create(
            seller=seller, amount_cents=100, status=InvoiceStatus.CANCELLED
        )
        # Счета без размещений — не у рабочего продукта: список всех продуктов (ADR-057).
        admin_client.post(
            f"{LIST_URL}?product=all",
            {"action": "mark_paid_action", "_selected_action": [issued.pk, cancelled.pk]},
        )
        issued.refresh_from_db()
        cancelled.refresh_from_db()
        assert issued.status == InvoiceStatus.PAID
        assert issued.paid_on == timezone.localdate()
        assert issued.paid_by is not None
        assert cancelled.status == InvoiceStatus.CANCELLED

    def test_list_shows_sites_and_pay_link(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        placement = _placement(product, "dev.to")
        invoice = Invoice.objects.create(
            seller=seller, amount_cents=18100, pay_url="https://www.paypal.com/invoice/p/#CL9"
        )
        InvoiceItem.objects.create(invoice=invoice, placement=placement, amount_cents=18100)
        page = admin_client.get(LIST_URL).content.decode()
        assert "dev.to" in page and "€181" in page
        assert 'href="https://www.paypal.com/invoice/p/#CL9" target="_blank"' in page
        filtered = admin_client.get(LIST_URL, {"status__exact": "paid"}).content.decode()
        assert "dev.to" not in filtered

    def test_list_queries_do_not_grow(
        self,
        admin_client: Client,
        product: Product,
        seller: Seller,
        django_assert_max_num_queries: Any,
    ) -> None:
        for n in range(15):
            invoice = Invoice.objects.create(seller=seller, amount_cents=100)
            InvoiceItem.objects.create(
                invoice=invoice, placement=_placement(product, f"s{n}.com"), amount_cents=100
            )
        with django_assert_max_num_queries(20):
            admin_client.get(LIST_URL)


class TestFromPlacements:
    def test_action_opens_form_with_selected(self, admin_client: Client, product: Product) -> None:
        first, second = _placement(product, "a.com"), _placement(product, "b.com")
        response = admin_client.post(
            PLACEMENTS_URL,
            {"action": "invoice_for_selected_action", "_selected_action": [second.pk, first.pk]},
        )
        assert response.status_code == 302
        assert response["Location"] == f"{ADD_URL}?placements={first.pk},{second.pk}"

    def test_form_prefilled(self, admin_client: Client, product: Product, seller: Seller) -> None:
        first = _placement(product, "a.com", seller=seller, price_paid_cents=18100)
        second = _placement(product, "b.com", seller=seller, price_paid_cents=9000)
        response = admin_client.get(ADD_URL, {"placements": f"{first.pk},{second.pk}"})
        form = response.context["adminform"].form
        assert form.initial["seller"] == seller.pk
        assert form.initial["amount_cents"] == 27100
        formset = response.context["inline_admin_formsets"][0].formset
        assert [f.initial for f in formset.forms] == [
            {"placement": first.pk, "amount_cents": 18100},
            {"placement": second.pk, "amount_cents": 9000},
        ]

    def test_form_without_sums_splits_evenly(self, admin_client: Client, product: Product) -> None:
        first, second = _placement(product, "a.com"), _placement(product, "b.com")
        response = admin_client.get(ADD_URL, {"placements": f"{first.pk},{second.pk}"})
        formset = response.context["inline_admin_formsets"][0].formset
        assert [f.initial["amount_cents"] for f in formset.forms] == [None, None]
        assert "seller" not in response.context["adminform"].form.initial

    def test_invoice_filter_and_paid_column(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        inside = _placement(product, "inside.com")
        _placement(product, "outside.com")
        invoice = Invoice.objects.create(seller=seller, amount_cents=18100)
        InvoiceItem.objects.create(invoice=invoice, placement=inside, amount_cents=18100)
        Placement.objects.filter(pk=inside.pk).update(price_paid_cents=18100)
        none = admin_client.get(PLACEMENTS_URL, {"invoice": "none"}).content.decode()
        assert "outside.com" in none and "inside.com" not in none
        issued = admin_client.get(PLACEMENTS_URL, {"invoice": "issued"}).content.decode()
        assert "inside.com" in issued and "outside.com" not in issued
        assert "счёт не оплачен" in issued


class TestPlacementForm:
    def test_paid_from_invoice_read_only(
        self, admin_client: Client, product: Product, seller: Seller
    ) -> None:
        placement = _placement(product, "dev.to")
        admin_client.post(ADD_URL, _form(seller, [(placement, "")], amount="181", number="12"))
        url = reverse("admin:placements_placement_change", args=[placement.pk])
        page = admin_client.get(url, headers=PARTIAL).content.decode()
        assert 'name="price_paid_cents"' not in page
        assert "из счёта, руками не правится" in page
        assert "счёт № 12 от" in page
        # Форма размещения без поля суммы её не стирает.
        data = {
            "status": "published",
            "placement_type": "",
            "seller": str(seller.pk),
            "employee": "",
            "collaborator_order_id": "",
            "ordered_at": "",
            "article_url": "",
            "published_at": "2026-09-30",
            "announce_on_homepage": "unknown",
            "clicks_from_homepage": "",
            "comment": "",
            "links-TOTAL_FORMS": "0",
            "links-INITIAL_FORMS": "0",
            "links-MIN_NUM_FORMS": "0",
            "links-MAX_NUM_FORMS": "1000",
        }
        response = admin_client.post(url, data, headers=PARTIAL)
        assert response.json()["saved"] is True
        assert _paid(placement) == (18100, "EUR")


def test_home_card_due(admin_client: Client, seller: Seller) -> None:
    Invoice.objects.create(seller=seller, amount_cents=18100)
    Invoice.objects.create(seller=seller, amount_cents=5000, currency="USD")
    Invoice.objects.create(seller=seller, amount_cents=999, status=InvoiceStatus.PAID)
    cards = admin_client.get(reverse("admin:index")).context["home_cards"]
    card = next(card for card in cards if card["name"] == "Счета")
    assert (card["value"], card["note"]) == ("2", "€181 + $50 к оплате")
    assert card["url"] == LIST_URL


class TestSeller:
    def test_card_and_list(self, admin_client: Client, product: Product, seller: Seller) -> None:
        due = Invoice.objects.create(seller=seller, amount_cents=18100, issued_on="2026-09-30")
        InvoiceItem.objects.create(
            invoice=due, placement=_placement(product, "dev.to"), amount_cents=18100
        )
        Invoice.objects.create(seller=seller, amount_cents=5000, currency="USD")
        Invoice.objects.create(
            seller=seller, amount_cents=20000, status=InvoiceStatus.PAID, paid_on="2026-09-01"
        )
        Invoice.objects.create(seller=seller, amount_cents=999, status=InvoiceStatus.CANCELLED)
        url = reverse("admin:sites_seller_change", args=[seller.pk])
        card = admin_client.get(url, headers=PARTIAL).content.decode()
        assert "К оплате: <b>€181 + $50</b>" in card
        assert '30.09.2026 · €181</a> <span class="seo-sub">dev.to</span>' in card
        assert "Заплачено по счетам: <b>€200</b>" in card
        assert f"?seller__id__exact={seller.pk}" in card
        assert f"{ADD_URL}?seller={seller.pk}" in card
        # Отменённый не считается.
        assert "€9.99" not in card
        page = admin_client.get(reverse("admin:sites_seller_changelist")).content.decode()
        assert "€181 + $50 · счетов: 2" in page
        assert "€200" in page

    def test_new_invoice_from_card_gets_seller(self, admin_client: Client, seller: Seller) -> None:
        response = admin_client.get(ADD_URL, {"seller": seller.pk})
        assert response.context["adminform"].form.initial["seller"] == str(seller.pk)

    def test_list_queries_do_not_grow(
        self, admin_client: Client, django_assert_max_num_queries: Any
    ) -> None:
        for n in range(15):
            seller = Seller.objects.create(name=f"seller {n}")
            Invoice.objects.create(seller=seller, amount_cents=100)
        with django_assert_max_num_queries(15):
            admin_client.get(reverse("admin:sites_seller_changelist"))


class TestExport:
    @pytest.fixture
    def bills(self, product: Product, seller: Seller) -> tuple[Invoice, Invoice]:
        seller.contacts = "Alexander@star5media.com"
        seller.save()
        dev = _placement(product, "dev.to", article_url="https://dev.to/ben_blog/licence")
        one = Invoice.objects.create(
            seller=seller,
            amount_cents=18100,
            pay_url="https://www.paypal.com/invoice/p/#CL9SJFQ7NL4VUWU7",
            status=InvoiceStatus.PAID,
            paid_on="2026-10-01",
            issued_on="2026-09-30",
        )
        InvoiceItem.objects.create(invoice=one, placement=dev, amount_cents=18100)
        batch = Invoice.objects.create(seller=seller, amount_cents=10000, issued_on="2026-10-02")
        for domain, cents in (("a.com", 3334), ("b.com", 3333), ("c.com", 3333)):
            item = InvoiceItem.objects.create(
                invoice=batch, placement=_placement(product, domain), amount_cents=cents
            )
            assert item.pk
        return one, batch

    def test_invoices_sheet(self, admin_client: Client, bills: tuple[Invoice, Invoice]) -> None:
        url = reverse("admin:placements_invoice_export", args=["xlsx"])
        response = admin_client.get(url, {"o": "2"})
        book = load_workbook(io.BytesIO(response.content))
        rows = list(book["Счета"].iter_rows(values_only=True))
        assert rows[0][:7] == (
            "Площадка",
            "Ссылка на статью",
            "Цена",
            "Цена за пачку",
            "Ссылка на оплату",
            "Вебмастер\\агентство",
            "Средняя цена за сайт",
        )
        by_site = {row[0]: row for row in rows[1:]}
        dev = by_site["dev.to"]
        assert dev[1:7] == (
            "https://dev.to/ben_blog/licence",
            181,
            None,
            "https://www.paypal.com/invoice/p/#CL9SJFQ7NL4VUWU7",
            "StarMedia",
            None,
        )
        assert by_site["a.com"][2:4] == (33.34, 100)
        assert by_site["b.com"][6] == pytest.approx(33.33)
        head = rows[0]
        assert dev[head.index("Статус")] == "Оплачен"
        assert dev[head.index("Контакты продавца")] == "Alexander@star5media.com"

    def test_monthly_report_has_invoice(
        self, admin_client: Client, bills: tuple[Invoice, Invoice]
    ) -> None:
        url = reverse("admin:placements_placement_export", args=["xlsx"])
        book = load_workbook(io.BytesIO(admin_client.get(url).content))
        rows: list[tuple[Any, ...]] = list(book["Размещения"].iter_rows(values_only=True))
        head = rows[0]
        by_site = {row[0]: row for row in rows[1:]}
        dev, a = by_site["dev.to"], by_site["a.com"]
        assert dev[head.index("Счёт")] == "счёт от 30.09.2026, StarMedia"
        assert dev[head.index("Счёт: статус")] == "Оплачен"
        assert dev[head.index("Счёт оплачен")].date() == dt.date(2026, 10, 1)
        assert dev[head.index("Счёт: ссылка на оплату")].startswith("https://www.paypal.com/")
        assert a[head.index("Счёт: статус")] == "Выставлен"
        assert a[head.index("Счёт оплачен")] is None
