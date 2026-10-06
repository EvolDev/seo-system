"""Отчёт за месяц: фильтр «опубликовано» и выгрузка «Размещений» (E1-06, ADR-053)."""

import csv
import datetime as dt
import io
from collections.abc import Callable, Iterator
from decimal import Decimal
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from django.contrib.auth.models import User
from django.db import connection
from django.test import Client, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from apps.placements.models import Placement, PlacementLink, PlacementStatus
from apps.sites.models import (
    ExchangeRate,
    PlacementType,
    Product,
    Seller,
    Site,
    SiteCountryMetric,
    SiteMetric,
    SiteNote,
    SitePrice,
)

pytestmark = [pytest.mark.django_db, pytest.mark.usefixtures("moscow")]

OfferFactory = Callable[..., SitePrice]
LIST_URL = reverse("admin:placements_placement_changelist")
MSK = ZoneInfo("Europe/Moscow")
SHEET_HEAD = [
    "Target",
    "Источник",
    "Комментарий к площадке",
    "URL статьи",
    "Индексация",
    "Статус",
    "Дата размещения",
    "Месяц",
    "Комментарий по размещению",
    "Анкор1",
    "Ссылка1",
    "Анкор2",
    "Ссылка2",
]


def _url(fmt: str = "xlsx") -> str:
    return reverse("admin:placements_placement_export", args=[fmt])


@pytest.fixture
def moscow() -> Iterator[None]:
    """Месяц — по времени проекта; в тестах — московскому, как по умолчанию."""
    with override_settings(TIME_ZONE="Europe/Moscow"):
        yield


@pytest.fixture
def base(offer: OfferFactory) -> dict[str, Any]:
    """Convertio: egg (сентябрь, Collaborator, евро), pod (30.09 23:30 по Москве, доллары,
    три ссылки), late (1 октября 00:30), wait (заявка без даты). Clideo: статья на egg."""
    convertio = Product.objects.create(name="Convertio", domain="convertio.co")
    clideo = Product.objects.create(name="Clideo", domain="clideo.com")
    egg = Site.objects.create(
        domain="eggradients.com",
        languages=["Английский"],
        site_type="Персональный блог",
        collaborator_url="https://collaborator.pro/ru/creator/article/view?id=490809",
        topics=["Культура и искусство"],
        marks_as_ad=False,
        link_type="dofollow",
        links_allowed=1,
    )
    SiteMetric.objects.create(
        site=egg,
        dr=55,
        organic_traffic=96653,
        top_geo="us",
        top_geo_traffic=60990,
        total_keywords=15707,
    )
    SiteCountryMetric.objects.create(site=egg, country="us", organic_traffic=60990)
    offer(egg, 54457, writing_cents=4084, announce_cents=0)
    SiteNote.objects.create(site=egg, body="На доработке пока что")
    pod, late, wait = (
        Site.objects.create(domain=domain) for domain in ("podchaser.com", "late.com", "wait.com")
    )
    linkhub = Seller.objects.create(name="LinkHub Media", currency="USD")
    ExchangeRate.objects.create(
        currency="USD", rate_date=dt.date(2026, 9, 30), rate=Decimal("1.1355")
    )
    ivan = User.objects.create_user("ivan", first_name="Иван", last_name="Петров")

    first = Placement.objects.create(
        site=egg,
        product=convertio,
        status=PlacementStatus.PLACED,
        placement_type=PlacementType.GUEST_POST,
        article_url="https://www.eggradients.com/blog/image-file-conversion",
        published_at=dt.datetime(2026, 9, 11, tzinfo=MSK),
        price_paid_cents=58182,
        currency="EUR",
        is_indexed=True,
        seller=Seller.collaborator(),
        employee=ivan,
        comment="=1+1 из чужого файла",
    )
    PlacementLink.objects.create(
        placement=first, link_index=2, anchor="Convertio", target_url="https://convertio.co/"
    )
    PlacementLink.objects.create(
        placement=first,
        link_index=1,
        anchor="image converter",
        target_url="https://convertio.co/image-converter/",
    )
    second = Placement.objects.create(
        site=pod,
        product=convertio,
        status=PlacementStatus.PLACED,
        placement_type=PlacementType.LINK_INSERTION,
        published_at=dt.datetime(2026, 9, 30, 23, 30, tzinfo=MSK),
        price_paid_cents=10000,
        currency="USD",
        seller=linkhub,
    )
    for index in (1, 2, 3):
        PlacementLink.objects.create(
            placement=second, link_index=index, anchor=f"a{index}", target_url=f"https://c/{index}"
        )
    Placement.objects.create(
        site=late,
        product=convertio,
        status=PlacementStatus.PLACED,
        published_at=dt.datetime(2026, 10, 1, 0, 30, tzinfo=MSK),
    )
    Placement.objects.create(site=wait, product=convertio, status=PlacementStatus.ORDERED)
    Placement.objects.create(
        site=egg,
        product=clideo,
        status=PlacementStatus.PLACED,
        article_url="https://www.eggradients.com/blog/top-7-color-perfection",
        published_at=dt.datetime(2026, 8, 1, tzinfo=MSK),
    )
    return {"convertio": convertio, "clideo": clideo}


def _book(client: Client, **params: str) -> Any:
    response = client.get(_url(), params)
    assert response.status_code == 200
    return response, load_workbook(io.BytesIO(response.content))


def _rows(sheet: Any) -> list[dict[str, Any]]:
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    return [dict(zip(values[0], row, strict=True)) for row in values[1:]]


def _september(base: dict[str, Any]) -> dict[str, str]:
    return {"product__id__exact": str(base["convertio"].pk), "month": "2026-09"}


def test_month_filter_by_project_time(admin_client: Client, base: dict[str, Any]) -> None:
    page = admin_client.get(LIST_URL).content.decode()
    # Месяцы, где что-то опубликовано, свежие первыми, и «без даты».
    assert page.index("октябрь 2026") < page.index("сентябрь 2026") < page.index("август 2026")
    assert "без даты" in page

    def domains(**params: str) -> list[str]:
        content = admin_client.get(LIST_URL, params).content.decode()
        return [d for d in ("eggradients", "podchaser", "late.com", "wait.com") if d in content]

    # 30.09 23:30 по Москве — ещё сентябрь, 1.10 00:30 — уже октябрь.
    assert domains(month="2026-09", product__id__exact=str(base["convertio"].pk)) == [
        "eggradients",
        "podchaser",
    ]
    assert domains(month="2026-10") == ["late.com"]
    assert domains(month="none") == ["wait.com"]
    assert domains(month="мусор") == []


def test_month_report_like_sheet(admin_client: Client, base: dict[str, Any]) -> None:
    response, book = _book(admin_client, **_september(base))
    assert book.sheetnames == ["Размещения", "Итого"]
    sheet = book["Размещения"]
    header = [cell.value for cell in next(sheet.iter_rows())]
    # Колонки листа «Размещения» — в том же порядке; у pod три ссылки — третья пара.
    assert header[: len(SHEET_HEAD)] == SHEET_HEAD
    assert header[13:16] == ["Анкор3", "Ссылка3", "Пример статьи на Clideo"]
    assert header[16:22] == [
        "Organic / Traffic",
        "Top Geo",
        "Top Geo Traff",
        "US Traff",
        "DR",
        "Organic / Total Keywords",
    ]
    assert header.index("Итог цена") < header.index("Тип ссылки статья") < header.index("Продукт")
    assert "Пример статьи на Convertio" not in header

    rows = {row["Target"]: row for row in _rows(sheet)}
    egg, pod = rows["eggradients.com"], rows["podchaser.com"]
    assert egg["Target"] == "eggradients.com"
    assert egg["Источник"] == "Collaborator"
    assert egg["Комментарий к площадке"] == "На доработке пока что"
    assert egg["Индексация"] == "Да"
    assert egg["Статус"] == "Размещено"
    assert egg["Дата размещения"] == dt.datetime(2026, 9, 11)
    assert egg["Месяц"] == dt.datetime(2026, 9, 1)
    # Текст из файла не стал формулой.
    assert egg["Комментарий по размещению"] == "=1+1 из чужого файла"
    assert (egg["Анкор1"], egg["Ссылка1"]) == (
        "image converter",
        "https://convertio.co/image-converter/",
    )
    assert (egg["Анкор2"], egg["Анкор3"]) == ("Convertio", None)
    assert (
        egg["Пример статьи на Clideo"] == "https://www.eggradients.com/blog/top-7-color-perfection"
    )
    assert (egg["DR"], egg["US Traff"], egg["Top Geo"]) == (55, 60990, "us")
    assert egg["Тип ссылки"] == "публикация"
    assert egg["Цена размещения статья, EUR"] == 544.57
    assert egg["Цена анонса статья, EUR"] == 0
    assert egg["Цена написания статья, EUR"] == 40.84
    assert egg["Итог цена"] == 581.82
    assert egg["Языки сайта"] == "Английский"
    assert egg["Пометка о рекламе статья"] == "Нет"
    assert egg["Тип ссылки статья"] == "dofollow"
    assert egg["Сотрудник"] == "Иван Петров"
    assert egg["Продукт"] == "Convertio"
    assert egg["Итог цена, EUR"] == 581.82

    assert pod["Анкор3"] == "a3"
    assert pod["Тип ссылки"] == "вставка ссылки"
    # $100 по курсу 1.1355 — €88.07.
    assert (pod["Итог цена"], pod["Валюта"], pod["Итог цена, EUR"]) == (100, "USD", 88.07)

    assert (
        "%D1%81%D0%B5%D0%BD%D1%82%D1%8F%D0%B1%D1%80%D1%8C%202026.xlsx"
        in (response["Content-Disposition"])
    )  # «Размещения Convertio сентябрь 2026.xlsx»


def test_totals_sheet(admin_client: Client, base: dict[str, Any]) -> None:
    _, book = _book(admin_client, **_september(base))
    totals = [[cell.value for cell in row] for row in book["Итого"].iter_rows()]
    assert totals[0] == ["Группа", "Значение", "Размещений", "Итог цена, EUR", "Без суммы"]
    assert totals[1] == ["Всего", None, 2, 669.89, 0]
    # Один статус и один продукт — строк по ним нет; продавцы — по сумме.
    assert totals[2:] == [
        ["Источник", "Collaborator", 1, 581.82, 0],
        ["Источник", "LinkHub Media", 1, 88.07, 0],
        ["Тип ссылки", "публикация", 1, 581.82, 0],
        ["Тип ссылки", "вставка ссылки", 1, 88.07, 0],
    ]


def test_working_product_and_statuses(admin_client: Client, base: dict[str, Any]) -> None:
    # Выгрузка — под рабочий продукт: строки списка его, «Все» их не добавляет (ADR-063).
    _, book = _book(admin_client, product__id__exact="all")
    header = [cell.value for cell in next(book["Размещения"].iter_rows())]
    # Выгрузка одного продукта: свой столбец примера не нужен, чужой остаётся.
    assert "Пример статьи на Clideo" in header
    assert "Пример статьи на Convertio" not in header
    rows = {(row["Target"], row["Продукт"]): row for row in _rows(book["Размещения"])}
    assert len(rows) == 4
    assert ("eggradients.com", "Clideo") not in rows
    # Статья Convertio видит статью Clideo на той же площадке.
    convertio = rows[("eggradients.com", "Convertio")]
    assert convertio["Пример статьи на Clideo"].endswith("top-7-color-perfection")
    totals = [[cell.value for cell in row] for row in book["Итого"].iter_rows()]
    assert totals[1] == ["Всего", None, 4, 669.89, 2]
    groups = {(row[0], row[1]): row[2:] for row in totals[2:]}
    assert ("Продукт", "Convertio") not in groups
    assert groups[("Статус", "Заявка отправлена")] == [1, 0, 1]
    assert groups[("Источник", "не указан")] == [2, 0, 2]


def test_csv_first_sheet_only(admin_client: Client, base: dict[str, Any]) -> None:
    response = admin_client.get(_url("csv"), _september(base))
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][: len(SHEET_HEAD)] == SHEET_HEAD
    assert len(rows) == 3
    egg = next(
        dict(zip(rows[0], row, strict=True)) for row in rows[1:] if row[0] == "eggradients.com"
    )
    assert (egg["Месяц"], egg["Итог цена"], egg["Индексация"]) == ("сентябрь 2026", "581,82", "Да")
    assert egg["Комментарий по размещению"] == "'=1+1 из чужого файла"


def test_buttons_and_queries(admin_client: Client, base: dict[str, Any]) -> None:
    page = admin_client.get(LIST_URL, {"month": "2026-09"}).content.decode()
    assert f'href="{_url()}?month=2026-09"' in page
    assert f'data-csv="{_url("csv")}?month=2026-09"' in page
    for group in ("Размещение", "Площадка на сегодня", "Сверх листа таблицы"):
        assert f"<legend>{group}</legend>" in page
    assert "Анкоры и ссылки" in page
    with CaptureQueriesContext(connection) as small:
        admin_client.get(_url(), {"month": "2026-10"})
    with CaptureQueriesContext(connection) as big:
        admin_client.get(_url())
    assert len(big) == len(small)


def test_file_name_without_month(admin_client: Client, base: dict[str, Any]) -> None:
    response = admin_client.get(_url(), {"month": "none"})
    assert "%D0%B1%D0%B5%D0%B7%20%D0%B4%D0%B0%D1%82%D1%8B.xlsx" in response["Content-Disposition"]


def test_chosen_columns_picked_rows_and_totals(admin_client: Client, base: dict[str, Any]) -> None:
    """Галочки «Анкоры и ссылки» и «Итог цена»; отмечено одно размещение — в файле оно,
    «Итого» — по нему."""
    egg = Placement.objects.get(site__domain="eggradients.com", product=base["convertio"])
    response = admin_client.post(
        _url() + "?month=2026-09",
        {"columns": ["paid", "links", "target"], "ids": [str(egg.pk)]},
    )
    book = load_workbook(io.BytesIO(response.content))
    rows = [[cell.value for cell in row] for row in book["Размещения"].iter_rows()]
    # Порядок листа: Target, пары ссылок (у egg две — колонок две пары), Итог цена.
    assert rows[0] == ["Target", "Анкор1", "Ссылка1", "Анкор2", "Ссылка2", "Итог цена"]
    assert rows[1][0] == "eggradients.com" and rows[1][-1] == 581.82
    assert len(rows) == 2
    totals = [[cell.value for cell in row] for row in book["Итого"].iter_rows()]
    assert totals[1] == ["Всего", None, 1, 581.82, 0]
    page = admin_client.get(LIST_URL).content.decode()
    assert 'name="columns" value="dr">' in page  # снятая галочка запомнилась
