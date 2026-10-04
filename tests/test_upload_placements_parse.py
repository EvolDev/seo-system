"""Файл размещений: разметка колонок и разбор строк (E1-09, ADR-051).

Файлы — синтетические, по образцу `import_examples/clideo_через_барыг.xlsx`
(сам пример в git не идёт): те же заголовки, значения как в настоящих строках.
"""

import datetime as dt
from collections.abc import Mapping, Sequence
from pathlib import Path

import pytest
from openpyxl import Workbook

from apps.placements.models import PlacementStatus
from apps.sites.models import PlacementType, Product, Upload, UploadKind, UploadStatus
from apps.sites.uploads import placement_records as records
from apps.sites.uploads.columns import Confidence
from apps.sites.uploads.files import Table, read_table
from apps.sites.uploads.placement_columns import (
    PField,
    build_mapping,
    guess_columns,
    looks_like_placements_header,
    remembered_mapping,
    validate_mapping,
)
from apps.sites.uploads.placement_records import IndexMark, LinkData, parse_placements

HEADERS = [
    "Target", "Person", "Source", "URL статьи", "Индексация", "Статус", "Дата размещения",
    "Месяц", "Комментарий", "Анкор1", "Ссылка1", "Анкор2", "Ссылка2", "Приложение", "Traffic",
    "Top Geo", "Top Geo Traf", "US Traf", "DR", "Organic / Total Keywords", "Цена размещ",
    "Цена анонса", "Цена написания", "Итог цена", "Кол-во ссылок", "Пометка о рекламе статья",
    "Особые тематики", "Языки сайта", "Тип сайта", "Collaborator URL", "Тематика",
    "Есть ли ссылка", "Индексация 02.09.26",
]  # fmt: skip
COLUMN = {header: index for index, header in enumerate(HEADERS)}
CLIDEO = "clideo.com"


def _table(tmp_path: Path, rows: Sequence[Sequence[object]], name: str = "clideo.xlsx") -> Table:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(list(row))
    path = tmp_path / name
    book.save(path)
    return read_table(path, is_header=looks_like_placements_header)


def _row_of(cells: Mapping[str, object]) -> list[object]:
    row: list[object] = [None] * len(HEADERS)
    for header, value in cells.items():
        row[COLUMN[header]] = value
    return row


TOMSGUIDE = {
    "Target": "tomsguide.com",
    "Person": "Evgeniy",
    "Source": "Athena Smith",
    "URL статьи": "https://www.tomsguide.com/computing/how-ai-tools-change-content",
    "Индексация": "да",
    "Статус": "Размещено",
    "Дата размещения": "14.09.2026",
    "Месяц": 46291.0,
    "Комментарий": "Link insert",
    "Анкор1": "video editor",
    "Ссылка1": "https://clideo.com/video-editor",
    "Анкор2": "Clideo",
    "Ссылка2": "Clideo",
    "Traffic": "4 086 679",
    "Top Geo": "us",
    "DR": 87.0,
    "Итог цена": "311,81",
    "Индексация 02.09.26": "#Н/Д",
}


def _mapping(table: Table) -> dict[str, PField]:
    return {key: guess.field for key, guess in guess_columns(table.columns).items()}


def _parse(tmp_path: Path, *rows: Mapping[str, object]) -> records.ParsedPlacements:
    table = _table(tmp_path, [_row_of(row) for row in rows])
    return parse_placements(table, _mapping(table), product_domain=CLIDEO)


class TestGuess:
    def test_clideo_file_headers(self, tmp_path: Path) -> None:
        table = _table(tmp_path, [_row_of(TOMSGUIDE)])
        fields = _mapping(table)
        assert fields["target"] == PField.DOMAIN
        assert fields["person"] == PField.EMPLOYEE
        assert fields["source"] == PField.SELLER
        assert fields["url статьи"] == PField.ARTICLE_URL
        assert fields["индексация"] == PField.INDEXED
        assert fields["индексация 02.09.26"] == PField.INDEXED
        assert fields["статус"] == PField.STATUS
        assert fields["дата размещения"] == PField.PUBLISHED
        assert fields["комментарий"] == PField.NOTE
        assert fields["анкор1"] == PField.ANCHOR_1
        assert fields["ссылка1"] == PField.LINK_1
        assert fields["анкор2"] == PField.ANCHOR_2
        assert fields["ссылка2"] == PField.LINK_2
        assert fields["цена размещ"] == PField.PRICE
        assert fields["цена анонса"] == PField.ANNOUNCE
        assert fields["цена написания"] == PField.WRITING
        assert fields["итог цена"] == PField.PAID
        # Метрики, месяц, приложение, карточка Collaborator — в прочие данные.
        for key in ("traffic", "top geo", "us traf", "dr", "месяц", "приложение", "тематика"):
            assert fields[key] == PField.EXTRA, key

    def test_source_is_asked_with_a_hint(self, tmp_path: Path) -> None:
        guess = guess_columns(_table(tmp_path, [_row_of(TOMSGUIDE)]).columns)["source"]
        assert guess.confidence == Confidence.LIKELY
        assert "через кого" in guess.hint

    def test_url_next_to_domain_is_the_article(self, tmp_path: Path) -> None:
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append(["Site", "URL", "Anchor", "Link", "Price"])
        sheet.append(["a.com", "https://a.com/post", "clideo", "https://clideo.com/", "120"])
        path = tmp_path / "list.xlsx"
        book.save(path)
        table = read_table(path, is_header=looks_like_placements_header)
        assert _mapping(table) == {
            "site": PField.DOMAIN,
            "url": PField.ARTICLE_URL,
            "anchor": PField.ANCHOR_1,
            "link": PField.LINK_1,
            "price": PField.PAID,
        }


class TestMappingMemory:
    def test_remembered_columns_are_not_asked(self, tmp_path: Path) -> None:
        table = _table(tmp_path, [_row_of(TOMSGUIDE)])
        remembered = {"target": "domain", "source": "seller"}
        guesses, questions = build_mapping(table.columns, remembered)
        assert guesses["source"].confidence == Confidence.REMEMBERED
        assert "source" not in questions
        assert "person" in questions
        # Пустая колонка — ни поля, ни вопроса.
        assert "тип сайта" not in questions

    @pytest.mark.django_db
    def test_memory_comes_from_confirmed_uploads(self) -> None:
        product = Product.objects.create(name="Clideo", domain=CLIDEO)

        def upload(mapping: dict[str, str], status: UploadStatus) -> None:
            Upload.objects.create(
                kind=UploadKind.PLACEMENTS,
                product=product,
                prices_date=dt.date(2026, 10, 4),
                file_name="x.xlsx",
                file_path="x.xlsx",
                file_sha256="0" * 64,
                mapping=mapping,
                status=status,
            )

        upload({"source": "seller", "цена": "price_paid"}, UploadStatus.DONE)
        upload({"source": "extra"}, UploadStatus.CHECKED)
        # Не подтверждена человеком — не в памяти.
        upload({"цена": "placement_price"}, UploadStatus.NEW)
        assert remembered_mapping() == {"source": "extra", "цена": "price_paid"}


class TestValidate:
    def test_domain_is_required(self, tmp_path: Path) -> None:
        table = _table(tmp_path, [_row_of(TOMSGUIDE)])
        mapping = {**_mapping(table), "target": PField.EXTRA}
        assert validate_mapping(mapping, table.columns) == [
            "Укажите колонку с площадкой — доменом или адресом сайта."
        ]

    def test_anchor_needs_its_link(self, tmp_path: Path) -> None:
        table = _table(tmp_path, [_row_of(TOMSGUIDE)])
        mapping = {**_mapping(table), "ссылка2": PField.EXTRA}
        assert validate_mapping(mapping, table.columns) == [
            "«Анкор 2» без «Ссылка 2 — куда ведёт»: ссылка без адреса не записывается."
        ]

    def test_two_indexation_columns_are_fine_two_dates_are_not(self, tmp_path: Path) -> None:
        table = _table(tmp_path, [_row_of(TOMSGUIDE)])
        assert validate_mapping(_mapping(table), table.columns) == []
        mapping = {**_mapping(table), "месяц": PField.PUBLISHED}
        assert validate_mapping(mapping, table.columns) == [
            "«Дата размещения» выбрано у нескольких колонок: «Дата размещения», «Месяц»."
        ]


class TestParse:
    def test_tomsguide_row(self, tmp_path: Path) -> None:
        parsed = _parse(tmp_path, TOMSGUIDE)
        [record] = parsed.records
        assert record.domain == "tomsguide.com"
        assert record.status == PlacementStatus.PUBLISHED
        assert record.published_on == dt.date(2026, 9, 14)
        assert (record.seller, record.employee) == ("Athena Smith", "Evgeniy")
        assert record.paid_cents == 31181
        assert record.offer_cents == 31181
        # «Link insert» — тип размещения, а не заметка.
        assert record.placement_type == PlacementType.LINK_INSERTION
        assert record.note is None
        assert record.links == [LinkData(1, "video editor", "https://clideo.com/video-editor")]
        # «#Н/Д» — пустая отметка; «да» без даты в заголовке — на дату файла.
        assert record.marks == [IndexMark(None, True, "Индексация")]
        assert record.extra["Traffic"] == "4 086 679"
        assert record.extra["Месяц"] == "46291"
        assert "Индексация 02.09.26" not in record.extra
        [issue] = parsed.issues[records.LINK_NOT_URL]
        assert issue.message == "ссылка 2: вместо адреса «Clideo» — ссылка не записана"

    def test_swapped_anchor_and_link_are_fixed(self, tmp_path: Path) -> None:
        row = {
            "Target": "a.com",
            "Анкор1": "https://clideo.com/compress-video",
            "Ссылка1": "video compressor",
            "Анкор2": "clideo.com",
            "Ссылка2": "Clideo",
        }
        parsed = _parse(tmp_path, row)
        [record] = parsed.records
        assert record.links == [
            LinkData(1, "video compressor", "https://clideo.com/compress-video"),
            LinkData(2, "Clideo", "https://clideo.com"),
        ]
        assert len(parsed.issues[records.LINK_SWAPPED]) == 2

    def test_app_store_link_is_kept_and_reported(self, tmp_path: Path) -> None:
        app = "https://apps.apple.com/us/app/clideo-video-editor/id1552262611"
        parsed = _parse(tmp_path, {"Target": "a.com", "Анкор1": "App Store", "Ссылка1": app})
        assert parsed.records[0].links == [LinkData(1, "App Store", app)]
        assert parsed.issues[records.LINK_FOREIGN][0].message.endswith(app)

    def test_unclear_date_is_left_empty(self, tmp_path: Path) -> None:
        parsed = _parse(tmp_path, {"Target": "a.com", "Дата размещения": "1.23.2026"})
        assert parsed.records[0].published_on is None
        [issue] = parsed.issues[records.BAD_DATE]
        assert issue.message == "дата «1.23.2026» не разобрана — оставлена пустой"

    def test_two_indexation_marks(self, tmp_path: Path) -> None:
        row = {"Target": "a.com", "Индексация": "да", "Индексация 02.09.26": "нет"}
        assert _parse(tmp_path, row).records[0].marks == [
            IndexMark(None, True, "Индексация"),
            IndexMark(dt.date(2026, 9, 2), False, "Индексация 02.09.26"),
        ]

    def test_domain_only_row(self, tmp_path: Path) -> None:
        [record] = _parse(tmp_path, {"Target": "https://www.b.com/"}).records
        assert record.domain == "b.com"
        assert record.status is None
        assert (record.article_url, record.links, record.marks) == (None, [], [])

    def test_bad_rows_are_errors(self, tmp_path: Path) -> None:
        parsed = _parse(
            tmp_path,
            {"Target": "a.com", "Статус": "Отказали"},
            {"Target": "b.com", "URL статьи": "https://other.com/post"},
            {"Target": "c.com", "Итог цена": "договорная"},
        )
        assert parsed.records == []
        messages = [error.message for error in parsed.errors]
        assert messages[0] == "«Статус»: статус «Отказали» не распознан"
        assert messages[1] == "«URL статьи»: «https://other.com/post» — не адрес статьи на b.com"
        assert messages[2].startswith("«Итог цена»: не число")

    def test_number_cells_and_excel_dates(self, tmp_path: Path) -> None:
        # Число без разделителя разрядов — как есть: большую цену проверит порог плана.
        row = {"Target": "a.com", "Итог цена": 241258.0, "Дата размещения": 46291}
        [record] = _parse(tmp_path, row).records
        assert record.paid_cents == 24_125_800
        assert record.published_on == dt.date(2026, 9, 26)

    def test_three_digits_after_the_dot_are_cents(self, tmp_path: Path) -> None:
        # В таблице коллеги «197.393»: точка ушла в разделитель разрядов, в xlsx —
        # 197 393 в формате «#,##0». После точки — не больше двух цифр: 197,39.
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append(HEADERS)
        sheet.append(
            _row_of(
                {
                    "Target": "a.com",
                    "Итог цена": 197393,
                    "Цена размещ": 4330755,
                    "Цена написания": 13.16,
                    "Traffic": 4086679,
                }
            )
        )
        for header in ("Итог цена", "Цена размещ", "Traffic"):
            sheet.cell(row=2, column=COLUMN[header] + 1).number_format = "#,##0"
        path = tmp_path / "grouped.xlsx"
        book.save(path)
        table = read_table(path, is_header=looks_like_placements_header)
        parsed = parse_placements(table, _mapping(table), product_domain=CLIDEO)
        [record] = parsed.records
        assert (record.paid_cents, record.price_cents, record.writing_cents) == (
            19739,
            433076,
            1316,
        )
        # Метрика в том же формате — обычное число, в прочие данные как есть.
        assert record.extra["Traffic"] == "4086679"
        assert [issue.message for issue in parsed.issues[records.PRICE_FIXED]] == [
            "«Итог цена»: 197.393 → 197,39",
            "«Цена размещ»: 4.330.755 → 4330,76",
        ]

    def test_duplicate_rows(self, tmp_path: Path) -> None:
        parsed = _parse(
            tmp_path,
            TOMSGUIDE,
            {
                **TOMSGUIDE,
                "URL статьи": "https://tomsguide.com/computing/how-ai-tools-change-content/",
            },
            {**TOMSGUIDE, "URL статьи": "https://tomsguide.com/other"},
        )
        assert [r.article_url for r in parsed.records] == [
            TOMSGUIDE["URL статьи"],
            "https://tomsguide.com/other",
        ]
        assert parsed.duplicates == [("tomsguide.com", (2, 3))]


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        ("Индексация 02.09.26", dt.date(2026, 9, 2)),
        ("Индексация 2.9.2026", dt.date(2026, 9, 2)),
        ("Индексация", None),
        ("Индексация 31.02.26", None),
    ],
)
def test_header_day(header: str, expected: dt.date | None) -> None:
    assert records.header_day(header) == expected
