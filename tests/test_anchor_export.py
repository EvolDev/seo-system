"""Анкоры в буфер и Excel: фильтры, выбор всех страниц, формат загрузки."""

import datetime as dt
import io
from pathlib import Path
from typing import Any

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from apps.keywords.models import AnchorType, Keyword
from apps.sites.models import Product, UploadKind
from apps.sites.uploads import service
from apps.workspace.products import choose_product

pytestmark = pytest.mark.django_db
PAGE = reverse("admin:keywords_keywordcoverage_changelist")
COPY = reverse("admin:keywords_keywordcoverage_copy")
EXCEL = reverse("admin:keywords_keywordcoverage_export")


@pytest.fixture
def product(admin_user: User) -> Product:
    product = Product.objects.create(name="Convertio", domain="convertio.co")
    choose_product(admin_user, product)
    Keyword.objects.bulk_create(
        Keyword(
            product=product,
            keyword=f"audio {i:03}",
            target_url="https://convertio.co/wav-mp3/",
            tool="Audio",
            volume=i,
            global_volume=i * 2,
        )
        for i in range(548)
    )
    Keyword.objects.create(
        product=product,
        keyword="Convertio",
        target_url="https://convertio.co/",
        anchor_type=AnchorType.BRANDED,
        tool="Main",
    )
    Keyword.objects.create(
        product=product,
        keyword="retired",
        target_url="https://convertio.co/",
        is_active=False,
        volume=None,
        global_volume=None,
    )
    other = Product.objects.create(name="Other", domain="other.example")
    Keyword.objects.create(product=other, keyword="foreign", target_url="https://other.example/")
    return product


def rows(client: Client, **params: Any) -> list[list[Any]]:
    response = client.get(EXCEL, params)
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Анкоры"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:D{sheet.max_row}"
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_copy_across_pages_and_filter_order(admin_client: Client, product: Product) -> None:
    response = admin_client.post(
        COPY + f"?product={product.pk}&tool=Audio&per_page=500&p=2&o=-1", {"select_across": "1"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["count"] == 548
    assert data["text"].splitlines() == [f"audio {i:03}" for i in reversed(range(548))]


def test_copy_selected_inside_filters(admin_client: Client, product: Product) -> None:
    a, b = Keyword.objects.filter(product=product, keyword__startswith="audio").order_by("id")[:2]
    outside = Keyword.objects.get(keyword="foreign")
    result = admin_client.post(
        COPY + f"?product={product.pk}&q=audio", {"ids": [b.pk, outside.pk], "select_across": "0"}
    )
    assert result.json() == {"text": b.keyword, "count": 1}
    assert a.keyword not in result.json()["text"]
    empty = admin_client.post(COPY + f"?product={product.pk}&q=absent", {"ids": [a.pk]})
    assert empty.json() == {"text": "", "count": 0}


@pytest.mark.parametrize("data", [{}, {"ids": ["not-an-id"]}])
def test_copy_needs_valid_selection(
    admin_client: Client, product: Product, data: dict[str, Any]
) -> None:
    assert admin_client.post(COPY + f"?product={product.pk}", data).status_code == 400


def test_excel_all_pages_four_columns_and_full_urls(admin_client: Client, product: Product) -> None:
    data = rows(admin_client, product=product.pk, tool="Audio", per_page="500", p="2", o="1")
    assert data[0] == ["Keyword", "URL", "Volume", "Global Volume"]
    assert len(data) == 549
    assert data[1] == ["audio 000", "https://convertio.co/wav-mp3/", 0, 0]
    assert data[-1] == ["audio 547", "https://convertio.co/wav-mp3/", 547, 1094]


@pytest.mark.parametrize("activity,expected", [("active", 549), ("inactive", 1), ("all", 550)])
def test_activity_filter_both_transfers(
    admin_client: Client, product: Product, activity: str, expected: int
) -> None:
    params = f"?product={product.pk}&activity={activity}"
    assert admin_client.post(COPY + params, {"select_across": "1"}).json()["count"] == expected
    assert len(rows(admin_client, product=product.pk, activity=activity)) == expected + 1


def test_search_and_naked_filter(admin_client: Client, product: Product) -> None:
    assert rows(admin_client, product=product.pk, kind="naked")[1:] == [
        ["Convertio", "https://convertio.co/", None, None]
    ]
    assert rows(admin_client, product=product.pk, q="audio 547")[1:] == [
        ["audio 547", "https://convertio.co/wav-mp3/", 547, 1094]
    ]
    assert len(rows(admin_client, product=product.pk, q="not found")) == 1


def test_excel_formula_is_plain_text(admin_client: Client, product: Product) -> None:
    Keyword.objects.create(product=product, keyword="=1+1", target_url="https://convertio.co/")
    response = admin_client.get(EXCEL, {"product": str(product.pk), "q": "=1+1"})
    sheet = load_workbook(io.BytesIO(response.content))["Анкоры"]
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"


def test_controls_and_methods(admin_client: Client, product: Product) -> None:
    html = admin_client.get(PAGE, {"product": str(product.pk), "tool": "Audio"}).content.decode()
    assert "Скопировать анкоры" in html and "Экспорт в Excel" in html
    assert "data-copy-domains" not in html
    assert "tool=Audio" in html
    assert admin_client.get(COPY).status_code == 405
    assert admin_client.post(EXCEL).status_code == 405
    assert admin_client.get(EXCEL, {"bad_field": "x"}).status_code == 400


def test_permission_required(client: Client, product: Product) -> None:
    user = User.objects.create_user("anchor-export-reader", is_staff=True)
    client.force_login(user)
    assert client.get(EXCEL).status_code == 403
    assert client.post(COPY, {"select_across": "1"}).status_code == 403


def test_excel_recognized_by_anchor_upload(
    admin_client: Client, product: Product, settings: Any, tmp_path: Path
) -> None:
    settings.UPLOADS_DIR = tmp_path
    response = admin_client.get(EXCEL, {"product": str(product.pk), "q": "audio 547"})
    upload = service.create_upload(
        SimpleUploadedFile("anchors.xlsx", response.content),
        kind=UploadKind.ANCHORS,
        product=product,
        seller=None,
        prices_date=dt.date(2026, 10, 9),
    ).upload
    assert not service.needs_questions(upload), upload.columns
    assert {
        key: value for key, value in (upload.mapping or {}).items() if not key.startswith("__")
    } == {
        "keyword": "keyword",
        "url": "url",
        "volume": "volume",
        "global volume": "global_volume",
    }
