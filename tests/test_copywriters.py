"""Копирайтеры (ADR-072): справочник с карточкой, выбор в размещении, копирайт в «Экономике»."""

import datetime as dt
import io
from zoneinfo import ZoneInfo

import pytest
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from apps.content.models import Copywriter
from apps.placements import economics
from apps.placements.economics import Period
from apps.placements.models import Placement, PlacementStatus
from apps.sites.models import ExchangeRate, Product, Site

pytestmark = pytest.mark.django_db

MSK = ZoneInfo("Europe/Moscow")
PARTIAL = {"X-Seo-Partial": "1"}
TODAY = dt.date(2026, 10, 14)


@pytest.fixture
def product() -> Product:
    return Product.objects.create(name="Convertio", domain="convertio.co")


@pytest.fixture
def anna() -> Copywriter:
    return Copywriter.objects.create(name="Anna", rate_cents=4000)


def _placement(
    product: Product, domain: str, status: PlacementStatus, **fields: object
) -> Placement:
    return Placement.objects.create(
        site=Site.objects.create(domain=domain), product=product, status=status, **fields
    )


def _panel_form(**fields: str) -> dict[str, str]:
    """Форма панели размещения, как её отправляет браузер."""
    data = {
        "status": "writing",
        "placement_type": "",
        "seller": "",
        "employee": "",
        "collaborator_order_id": "",
        "ordered_at": "",
        "price_paid_cents": "",
        "currency": "EUR",
        "copywriter": "",
        "copywriter_cents": "",
        "copywriter_currency": "",
        "draft_url": "",
        "article_url": "",
        "published_at": "",
        "announce_on_homepage": "unknown",
        "clicks_from_homepage": "",
        "comment": "",
        "links-TOTAL_FORMS": "0",
        "links-INITIAL_FORMS": "0",
        "links-MIN_NUM_FORMS": "0",
        "links-MAX_NUM_FORMS": "1000",
    }
    data.update(fields)
    return data


class TestPlacement:
    def url(self, placement: Placement) -> str:
        return reverse("admin:placements_placement_change", args=[placement.pk])

    def test_rate_fills_empty_price(
        self, admin_client: Client, product: Product, anna: Copywriter
    ) -> None:
        placement = _placement(product, "a.com", PlacementStatus.WRITING)
        response = admin_client.post(
            self.url(placement),
            _panel_form(copywriter=str(anna.pk), draft_url="https://docs.google.com/d/1"),
            headers=PARTIAL,
        )
        assert response.json()["saved"] is True
        placement.refresh_from_db()
        assert placement.copywriter == anna
        assert (placement.copywriter_cents, placement.copywriter_currency) == (4000, "EUR")
        assert placement.draft_url == "https://docs.google.com/d/1"

    def test_typed_price_wins_and_survives_new_rate(
        self, admin_client: Client, product: Product, anna: Copywriter
    ) -> None:
        placement = _placement(product, "b.com", PlacementStatus.WRITING)
        admin_client.post(
            self.url(placement),
            _panel_form(copywriter=str(anna.pk), copywriter_cents="55", copywriter_currency="USD"),
            headers=PARTIAL,
        )
        anna.rate_cents = 9000
        anna.save()
        placement.refresh_from_db()
        assert (placement.copywriter_cents, placement.copywriter_currency) == (5500, "USD")

    def test_no_price_no_currency(self, admin_client: Client, product: Product) -> None:
        placement = _placement(product, "c.com", PlacementStatus.WRITING)
        admin_client.post(
            self.url(placement), _panel_form(copywriter_currency="EUR"), headers=PARTIAL
        )
        placement.refresh_from_db()
        assert (placement.copywriter_id, placement.copywriter_currency) == (None, None)

    def test_choices_and_rates(
        self, admin_client: Client, product: Product, anna: Copywriter
    ) -> None:
        gone = Copywriter.objects.create(name="Gone", rate_cents=3000, is_active=False)
        placement = _placement(product, "d.com", PlacementStatus.WRITING)
        html = admin_client.get(self.url(placement), headers=PARTIAL).content.decode()
        assert 'data-rate="40"' in html and 'data-currency="EUR"' in html
        assert ">Gone<" not in html
        # Назначенный раньше, а теперь не работающий, — остаётся выбранным.
        Placement.objects.filter(pk=placement.pk).update(copywriter=gone)
        html = admin_client.get(self.url(placement), headers=PARTIAL).content.decode()
        assert ">Gone<" in html


class TestCard:
    def test_list_and_card(self, admin_client: Client, product: Product, anna: Copywriter) -> None:
        ExchangeRate.objects.create(currency="USD", rate_date=TODAY, rate="1.25")
        _placement(
            product,
            "writing.com",
            PlacementStatus.WRITING,
            copywriter=anna,
            copywriter_cents=4000,
            copywriter_currency="EUR",
            draft_url="https://docs.google.com/d/draft",
        )
        _placement(
            product,
            "placed.com",
            PlacementStatus.PLACED,
            copywriter=anna,
            copywriter_cents=5000,
            copywriter_currency="USD",
            article_url="https://placed.com/post",
            published_at=dt.datetime(2026, 10, 2, 12, tzinfo=MSK),
        )
        _placement(product, "inwork.com", PlacementStatus.IN_WORK, copywriter=anna)
        listing = admin_client.get(reverse("admin:content_copywriter_changelist"))
        row = listing.context["cl"].result_list[0]
        assert (row.writing_total, row.placed_total) == (1, 1)
        assert round(row.earned_eur) == 4000  # $50 / 1.25
        card = admin_client.get(
            reverse("admin:content_copywriter_change", args=[anna.pk]), headers=PARTIAL
        ).content.decode()
        assert "writing.com" in card and "placed.com" in card
        assert "inwork.com" not in card  # только «Написание статьи» и «Размещено»
        assert "https://docs.google.com/d/draft" in card and "https://placed.com/post" in card
        assert "Пишет сейчас: <b>1</b>" in card
        assert card.index("writing.com") < card.index("placed.com")

    def test_menu(self, admin_client: Client) -> None:
        menu = admin_client.get(reverse("admin:index")).context["available_apps"]
        reference = next(group for group in menu if group["app_label"] == "reference")
        assert "Копирайтеры" in [str(model["name"]) for model in reference["models"]]

    def test_delete_clears_placements(
        self, admin_client: Client, product: Product, anna: Copywriter
    ) -> None:
        placement = _placement(product, "e.com", PlacementStatus.WRITING, copywriter=anna)
        response = admin_client.post(
            reverse("admin:content_copywriter_delete", args=[anna.pk]), {"post": "yes"}
        )
        assert response.status_code == 302
        placement.refresh_from_db()
        assert placement.copywriter_id is None


def test_economics_copy(product: Product, anna: Copywriter) -> None:
    boris = Copywriter.objects.create(name="Boris", rate_cents=3000)
    _placement(
        product,
        "p1.com",
        PlacementStatus.PLACED,
        published_at=dt.datetime(2026, 10, 2, 12, tzinfo=MSK),
        price_paid_cents=20000,
        copywriter=anna,
        copywriter_cents=4000,
        copywriter_currency="EUR",
    )
    _placement(
        product,
        "p2.com",
        PlacementStatus.WRITING,
        price_paid_cents=10000,
        copywriter=boris,
        copywriter_cents=3000,
        copywriter_currency="EUR",
    )
    _placement(product, "p3.com", PlacementStatus.WRITING, price_paid_cents=10000, copywriter=boris)
    october = Period(dt.date(2026, 10, 1), dt.date(2026, 11, 1))
    report = economics.report(october, [(product.pk, "Convertio")], product.pk, TODAY)
    total = report.total
    assert (total.copy_cents, total.copy_planned_cents) == (4000, 3000)
    assert total.with_copy_cents == 24000 and total.average_with_copy_cents == 24000
    # Копирайт не из бюджета: остаток считают без него.
    assert total.left_cents == 3_000_000 - 20000 - 20000
    writers = {
        group.name: (group.count, group.spent_cents, group.planned_cents)
        for group in report.by_copywriter
    }
    assert writers == {"Anna": (1, 4000, 0), "Boris": (2, 0, 3000)}
    assert report.no_copy_price == 1


def test_economics_screen_and_excel(
    admin_client: Client, product: Product, anna: Copywriter
) -> None:
    _placement(
        product,
        "p1.com",
        PlacementStatus.PLACED,
        published_at=dt.datetime.now(tz=MSK),
        price_paid_cents=20000,
        copywriter=anna,
        copywriter_cents=4000,
        copywriter_currency="EUR",
    )
    url = reverse("admin:economics")
    html = admin_client.get(url, {"product": str(product.pk)}).content.decode()
    assert "Копирайтеры" in html and "Anna" in html
    assert "Потрачено + копирайт" in html and "Колонки копирайта появятся" not in html
    assert 'data-econ-show="copy"' in html
    response = admin_client.get(url, {"product": str(product.pk), "export": "xlsx"})
    book = load_workbook(io.BytesIO(response.content))
    months = book["По месяцам"]
    headers = [cell.value for cell in months[1]]
    assert "Копирайт, €" in headers and "Потрачено + копирайт, €" in headers
    copy = months.cell(row=months.max_row - 1, column=headers.index("Копирайт, €") + 1).value
    assert copy == 40
    rows = book["Размещения"]
    assert rows.cell(row=2, column=13).value == "Anna"
