"""Запись выгрузки в файл: Excel и CSV для человека (E1-06, ADR-053)."""

import csv
import datetime as dt
import io
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.test import override_settings
from openpyxl import load_workbook

from config.export import Column, Format, Kind, Sheet, Value, csv_bytes, response, xlsx

COLUMNS = (
    Column("Домен", width=30),
    Column("DR", Kind.INT),
    Column("Цена", Kind.MONEY),
    Column("Доля серых, %", Kind.NUMBER),
    Column("Дата цены", Kind.DATE),
    Column("Пишем мы"),
    Column("Тематики"),
)

ROWS: tuple[tuple[Value, ...], ...] = (
    (
        "пример.укр",
        55,
        123456,
        Decimal("3.5"),
        dt.datetime(2026, 9, 30, 22, 30, tzinfo=dt.UTC),
        True,
        ["Бизнес и финансы", "СМИ"],
    ),
    ("b.com", None, None, None, dt.date(2026, 10, 1), False, []),
    ('=HYPERLINK("http://x","y")', 1, 5, 0, None, None, None),
)


def _sheet() -> Sheet:
    return Sheet("Площадки", COLUMNS, ROWS)


@override_settings(TIME_ZONE="Europe/Moscow")
def test_xlsx_cells_formats_and_header() -> None:
    book = load_workbook(io.BytesIO(xlsx([_sheet()])))
    sheet = book["Площадки"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:G4"
    assert sheet.column_dimensions["A"].width == 30
    rows = list(sheet.iter_rows())
    assert [cell.value for cell in rows[0]][:3] == ["Домен", "DR", "Цена"]
    assert all(cell.font.b for cell in rows[0])

    domain, dr, price, gray, day, we_write, topics = rows[1]
    assert domain.value == "пример.укр"
    assert dr.value == 55
    # Деньги — число с копейками, не центы и не текст.
    assert price.value == 1234.56
    assert price.number_format == "#,##0.00"
    assert gray.value == 3.5
    # Момент — день по времени проекта: 22:30 UTC — уже 1 октября в Москве.
    assert day.value == dt.datetime(2026, 10, 1)
    assert day.number_format == "DD.MM.YYYY"
    assert we_write.value == "Да"
    assert topics.value == "Бизнес и финансы, СМИ"

    assert [cell.value for cell in rows[2]] == [
        "b.com",
        None,
        None,
        None,
        dt.datetime(2026, 10, 1),
        "Нет",
        None,
    ]


def test_xlsx_text_never_becomes_formula() -> None:
    sheet = load_workbook(io.BytesIO(xlsx([_sheet()])))["Площадки"]
    cell = sheet["A4"]
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("http://x","y")'


def test_xlsx_several_sheets() -> None:
    total = Sheet("Итого", (Column("Размещений", Kind.INT),), [(3,)])
    book = load_workbook(io.BytesIO(xlsx([_sheet(), total])))
    assert book.sheetnames == ["Площадки", "Итого"]
    assert book["Итого"]["A2"].value == 3


@override_settings(TIME_ZONE="Europe/Moscow")
def test_csv_for_russian_excel() -> None:
    content = csv_bytes(_sheet())
    # BOM — по нему Excel узнаёт UTF-8, без него кириллица кракозябрами.
    assert content.startswith(b"\xef\xbb\xbf")
    text = content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert rows[0] == [
        "Домен",
        "DR",
        "Цена",
        "Доля серых, %",
        "Дата цены",
        "Пишем мы",
        "Тематики",
    ]
    assert rows[1] == [
        "пример.укр",
        "55",
        "1234,56",
        "3,50",
        "01.10.2026",
        "Да",
        "Бизнес и финансы, СМИ",
    ]
    assert rows[2] == ["b.com", "", "", "", "01.10.2026", "Нет", ""]
    # Формула из чужого файла — текстом, с апострофом.
    assert rows[3][0] == '\'=HYPERLINK("http://x","y")'
    assert rows[3][2] == "0,05"
    assert "\r\n" in text


def test_response_file_name_and_type() -> None:
    xlsx_response = response([_sheet()], "Площадки Convertio 04.10.2026", Format.XLSX)
    assert xlsx_response["Content-Type"].startswith("application/vnd.openxmlformats")
    disposition = xlsx_response["Content-Disposition"]
    assert disposition.startswith("attachment;")
    assert "filename*=utf-8''" in disposition
    assert disposition.endswith(".xlsx")

    csv_response = response([_sheet()], "Площадки", Format.CSV)
    assert csv_response["Content-Type"] == "text/csv; charset=utf-8"
    assert csv_response["Content-Disposition"].endswith(".csv")
    assert csv_response.content.startswith(b"\xef\xbb\xbf")


def test_day_in_project_time_zone() -> None:
    moment = dt.datetime(2026, 9, 30, 23, 0, tzinfo=ZoneInfo("Europe/Moscow"))
    with override_settings(TIME_ZONE="Europe/Moscow"):
        text = csv_bytes(Sheet("x", (Column("д", Kind.DATE),), [(moment,)])).decode("utf-8-sig")
    assert text.splitlines()[1] == "30.09.2026"


@override_settings(TIME_ZONE="Europe/Moscow")
def test_month_column() -> None:
    """Месяц, как в листе «Размещения»: в Excel — дата первого числа, в CSV — словами."""
    moment = dt.datetime(2026, 9, 30, 22, 30, tzinfo=dt.UTC)  # уже 1 октября в Москве
    sheet = Sheet("x", (Column("Месяц", Kind.MONTH),), [(moment,), (dt.date(2026, 9, 11),)])
    cells = [row[0] for row in load_workbook(io.BytesIO(xlsx([sheet])))["x"].iter_rows(min_row=2)]
    assert [cell.value for cell in cells] == [dt.datetime(2026, 10, 1), dt.datetime(2026, 9, 1)]
    assert cells[0].number_format == "MMMM YYYY"
    lines = csv_bytes(sheet).decode("utf-8-sig").splitlines()
    assert lines[1:] == ["октябрь 2026", "сентябрь 2026"]
