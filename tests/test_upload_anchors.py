"""Файл анкоров продукта — экран «Загрузки» (E3-05, ADR-059).

Книга — как таблица линкбилдинга: лист анкоров не первый, рядом листы долей.
"""

import datetime as dt
import io
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook

from apps.keywords.models import (
    AnchorType,
    CountryShare,
    Keyword,
    KeywordPosition,
    PageTypeShare,
)
from apps.placements.models import Placement, PlacementLink, PlacementStatus
from apps.sites.models import Product, Site, Upload, UploadKind, UploadStatus
from apps.sites.tasks import upload_check, upload_write
from apps.sites.uploads import anchors, service

pytestmark = pytest.mark.django_db

D = Decimal
HEADER = [
    "Keyword", "URL", "Volume", "Global Volume", "Links Placed", "Links Waiting",
    "Pos 22.09.26", "Pos 16.09.26", "Tool", "Type",
]  # fmt: skip
ROWS: list[list[Any]] = [
    ["convert", "https://convertio.co/", 42000.0, 363000.0, 5, 2, 7, 14, "Main", "Главная"],
    ["mp4 converter", "https://convertio.co/mp4-converter/", 9000, 50000, None, 1, 12, 101,
     "Video", "Video (xxx-yyy)"],
    ["mp3 to wav", "https://convertio.co/mp3-wav/", "n/a", None, None, None, None, 30,
     "Audio", "Audio (xxx-yyy)"],
]  # fmt: skip


def _book(rows: list[list[Any]] | None = None, *, shares: bool = True) -> bytes:
    book = Workbook()
    first = book.active
    assert first is not None
    first.title = "Размещения"
    first.append(["Target", "Source"])
    first.append(["example.com", "Someone"])
    sheet = book.create_sheet("Распределение анкоров")
    sheet.append(HEADER)
    for row in rows if rows is not None else ROWS:
        sheet.append(row)
    if shares:
        types = book.create_sheet("Распределение по типам страниц")
        types.append(["Тип страницы", "Общая целевая доля", "Текущее распределение размещенных",
                      "Прямой анкор", "Разбавленный анкор", "Безанкор"])  # fmt: skip
        types.append(["Главная", 0.15, 0.23, 0.5, 0.25, 0.25])
        types.append(["Video (xxx-yyy)", 0.25, 0.26, 0.85, 0.1, 0.05])
        types.append(["Audio (xxx-yyy)", 0.25, 0.11, 0.85, 0.1, 0.05])
        naked = book.create_sheet("Распределение безанкорки")
        naked.append(["Anchor text", "URL", "%"])
        naked.append(["Convertio", "https://convertio.co/", 56.86])
        naked.append(["convertio.co", "https://convertio.co/", 20.8])
        naked.append(["click here", "https://convertio.co/", 0.26])
        countries = book.create_sheet("Распределение по странам")
        countries.append(["Страна", "Целевая доля размещений"])
        countries.append(["США", 0.5])
        countries.append(["Бразилия", 0.3])
        countries.append(["Остальные", 0.2])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture(autouse=True)
def uploads_dir(settings: object, tmp_path: Path) -> None:
    settings.UPLOADS_DIR = tmp_path / "uploads"  # type: ignore[attr-defined]


@pytest.fixture
def convertio() -> Product:
    return Product.objects.create(name="Convertio", domain="convertio.co")


def _load(product: Product, data: bytes, day: dt.date = dt.date(2026, 10, 5)) -> Upload:
    # Тот же файл с той же датой — открывается прошлая загрузка; другая дата — новая проверка.
    file = SimpleUploadedFile("anchors.xlsx", data)
    upload = service.create_upload(
        file,
        kind=UploadKind.ANCHORS,
        seller=None,
        product=product,
        prices_date=day,
    ).upload
    assert upload.status == UploadStatus.NEW, upload.error
    assert not service.needs_questions(upload), upload.columns
    service.start(upload, UploadStatus.CHECKING)
    upload_check.delay(upload.pk)
    upload.refresh_from_db()
    assert upload.status == UploadStatus.CHECKED, upload.error
    return upload


def _write(upload: Upload) -> Upload:
    service.start(upload, UploadStatus.WRITING)
    upload_write.delay(upload.pk)
    upload.refresh_from_db()
    assert upload.status == UploadStatus.DONE, upload.error
    return upload


def test_columns_found_on_anchors_sheet(convertio: Product) -> None:
    upload = service.create_upload(
        SimpleUploadedFile("a.xlsx", _book()),
        kind=UploadKind.ANCHORS,
        seller=None,
        product=convertio,
        prices_date=dt.date(2026, 10, 5),
    ).upload
    fields = {c["header"]: c["field"] for c in upload.columns or []}
    assert fields["Keyword"] == "keyword"
    assert fields["URL"] == "url"
    assert fields["Pos 22.09.26"] == "position"
    assert fields["Links Placed"] == "skip"
    assert fields["Type"] == "page_type"


def test_first_load_creates_everything(convertio: Product) -> None:
    upload = _load(convertio, _book())
    summary = upload.summary or {}
    assert (summary["anchors"], summary["new_total"], summary["same"]) == (3, 3, 0)
    assert summary["dates"] == ["16.09.2026", "22.09.2026"]
    assert summary["positions_new"] == 5  # 101 — тоже снимок: «вне топ-100»
    assert any("n/a" in line for line in summary["issues"])
    _write(upload)
    convert = Keyword.objects.get(product=convertio, keyword="convert")
    assert (convert.volume, convert.global_volume, convert.anchor_type) == (
        42000,
        363000,
        AnchorType.EXACT,
    )
    assert convert.page_type == "Главная"
    mp4 = Keyword.objects.get(keyword="mp4 converter")
    positions = dict(
        KeywordPosition.objects.filter(keyword=mp4).values_list("checked_at", "position")
    )
    assert positions == {dt.date(2026, 9, 22): 12, dt.date(2026, 9, 16): None}
    shares = {s.page_type: s for s in PageTypeShare.objects.filter(product=convertio)}
    assert shares["Главная"].target_pct == D("15.00")
    assert shares["Главная"].naked_pct == D("25.00")
    assert shares["Audio (xxx-yyy)"].position == 3
    brand = Keyword.objects.get(keyword="Convertio")
    assert (brand.anchor_type, brand.share, brand.page_type) == (
        AnchorType.BRANDED,
        D("56.86"),
        "Главная",
    )
    assert Keyword.objects.get(keyword="convertio.co").anchor_type == AnchorType.URL
    assert Keyword.objects.get(keyword="click here").anchor_type == AnchorType.GENERIC
    countries = dict(
        CountryShare.objects.filter(product=convertio).values_list("country", "target_pct")
    )
    assert countries == {"us": D("50.00"), "br": D("30.00"), None: D("20.00")}


def test_second_load_of_same_file_changes_nothing(convertio: Product) -> None:
    _write(_load(convertio, _book()))
    again = _load(convertio, _book(), dt.date(2026, 10, 6))
    summary = again.summary or {}
    assert (summary["new_total"], summary["changed_total"], summary["same"]) == (0, 0, 3)
    assert (summary["positions_new"], summary["positions_changed"]) == (0, 0)
    assert (summary["naked_new"], summary["naked_changed"], summary["types_changed"]) == (0, 0, 0)
    assert summary["countries_changed"] == 0


def test_merge_keeps_what_file_lacks(convertio: Product) -> None:
    own = Keyword.objects.create(
        product=convertio, keyword="heic to jpg", target_url="https://convertio.co/heic-jpg/"
    )
    Keyword.objects.create(
        product=convertio, keyword="convert", target_url="https://convertio.co/old/", volume=1
    )
    rows = [
        ["convert", "https://convertio.co/", None, None, None, None, 7, None, "Main", "Главная"],
    ]
    upload = _load(convertio, _book(rows, shares=False))
    summary = upload.summary or {}
    assert summary["changed_total"] == 1
    changes = summary["changed"][0]["changes"]
    assert ["Куда ведёт", "https://convertio.co/old/", "https://convertio.co/"] in changes
    assert summary["missing"] == ["heic to jpg"]
    _write(upload)
    convert = Keyword.objects.get(keyword="convert")
    # Пустая ячейка объёма ничего не стёрла.
    assert (convert.target_url, convert.volume) == ("https://convertio.co/", 1)
    own.refresh_from_db()
    assert own.is_active


def test_naked_attaches_old_links(convertio: Product) -> None:
    placement = Placement.objects.create(
        site=Site.objects.create(domain="a.com"),
        product=convertio,
        status=PlacementStatus.PLACED,
    )
    link = PlacementLink.objects.create(
        placement=placement, anchor="convertio", target_url="https://convertio.co/"
    )
    upload = _write(_load(convertio, _book()))
    link.refresh_from_db()
    assert link.keyword is not None and link.keyword.keyword == "Convertio"
    assert link.anchor_type == AnchorType.BRANDED
    assert (upload.result or {})["counts"]["links"] == 1


def test_anchors_without_url_column_asked(convertio: Product) -> None:
    upload = service.create_upload(
        SimpleUploadedFile("a.csv", "Ключ,Частота\nconvert,100\n".encode()),
        kind=UploadKind.ANCHORS,
        seller=None,
        product=convertio,
        prices_date=dt.date(2026, 10, 5),
    ).upload
    errors = service.confirm_mapping(upload, {"ключ": "keyword", "частота": "volume"}, "")
    assert errors == ["Укажите колонку «куда ведёт» — адрес страницы продукта."]


@pytest.mark.parametrize(
    ("header", "date"),
    [("Pos 22.09.26", dt.date(2026, 9, 22)), ("Позиция 01.10.2026", dt.date(2026, 10, 1)),
     ("Pos", None)],
)  # fmt: skip
def test_position_date(header: str, date: dt.date | None) -> None:
    assert anchors.position_date(header) == date


def test_upload_form_preselects_anchors(admin_client: Client, convertio: Product) -> None:
    url = reverse("admin:sites_upload_add") + f"?kind=anchors&product={convertio.pk}"
    page = admin_client.get(url).content.decode()
    assert 'value="anchors" checked' in page or 'value="anchors"\n' in page or "checked" in page


def test_summary_page(admin_client: Client, convertio: Product) -> None:
    upload = _load(convertio, _book())
    page = admin_client.get(reverse("admin:sites_upload_summary", args=[upload.pk]))
    text = page.content.decode()
    assert "Анкоры <b>Convertio</b>" in text
    assert "Записать в базу" in text
    assert "mp4 converter" in text
