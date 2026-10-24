"""Загрузки и регионы списков анкоров (ADR-075): файл анкоров — в список своего
региона, лист книги — по пометке региона, ссылки файла размещений — к анкорам
региона площадки."""

import datetime as dt
import io
from pathlib import Path

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook

from apps.keywords.models import AnchorList, CountryShare, Keyword, KeywordPosition, PageTypeShare
from apps.placements.models import Placement
from apps.sites.models import Product, Site, Upload, UploadKind, UploadStatus
from apps.sites.tasks import upload_check, upload_write
from apps.sites.uploads import anchors, service, undo
from apps.workspace.products import choose_product

pytestmark = pytest.mark.django_db

HEADERS = ["Keyword", "URL", "Volume", "Pos 19.10.26", "Type"]


@pytest.fixture(autouse=True)
def uploads_dir(settings: object, tmp_path: Path) -> None:
    settings.UPLOADS_DIR = tmp_path / "uploads"  # type: ignore[attr-defined]


@pytest.fixture
def convertio() -> Product:
    return Product.objects.create(name="Convertio", domain="convertio.co")


@pytest.fixture
def us(convertio: Product) -> AnchorList:
    """Список США с анкорами — как после миграции: анкоры до регионов."""
    anchor_list = AnchorList.objects.create(product=convertio, country="us")
    Keyword.objects.create(
        anchor_list=anchor_list, keyword="convert", target_url="https://convertio.co/"
    )
    Keyword.objects.create(
        anchor_list=anchor_list, keyword="old us key", target_url="https://convertio.co/old/"
    )
    return anchor_list


def _book(*anchor_sheets: tuple[str, list[list[object]]], shares: bool = True) -> bytes:
    book = Workbook()
    for empty in list(book.worksheets):
        book.remove(empty)
    for title, rows in anchor_sheets:
        sheet = book.create_sheet(title)
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
    if shares:
        types = book.create_sheet("Распределение по типам страниц")
        types.append(["Тип страницы", "Общая целевая доля", "Прямой анкор", "Безанкор"])
        types.append(["Главная", 0.15, 0.5, 0.25])
        naked = book.create_sheet("Распределение безанкорки")
        naked.append(["Anchor text", "URL", "%"])
        naked.append(["Convertio", "https://convertio.co/", 56.86])
        countries = book.create_sheet("Распределение по странам")
        countries.append(["Страна", "Целевая доля размещений"])
        countries.append(["США", 0.5])
        countries.append(["Бразилия", 0.3])
        countries.append(["Остальные", 0.2])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


EN = (
    "Распределение анкоров EN",
    [
        ["convert", "https://convertio.co/", 42000, 5, "Главная"],
        ["mp4 to mp3", "https://convertio.co/mp4-mp3/", 9000, 12, "Video (xxx-yyy)"],
    ],
)
BR = (
    "Распределение анкоров BR",
    [
        ["convert", "https://convertio.co/pt/", 9100, 1, "Главная"],
        ["conversor de vídeo", "https://convertio.co/pt/video-converter/", 2600, 3, "Главная"],
    ],
)


def _create(
    product: Product, data: bytes, country: str, name: str = "book.xlsx", **options: object
) -> Upload:
    return service.create_upload(
        SimpleUploadedFile(name, data),
        kind=UploadKind.ANCHORS,
        seller=None,
        product=product,
        country=country,
        prices_date=dt.date(2026, 10, 24),
        **options,  # type: ignore[arg-type]
    ).upload


def _check(upload: Upload) -> Upload:
    assert upload.status == UploadStatus.NEW, upload.error
    service.start(upload, UploadStatus.CHECKING)
    upload_check.delay(upload.pk)
    upload.refresh_from_db()
    assert upload.status == UploadStatus.CHECKED, upload.error
    return upload


def _write(upload: Upload) -> Upload:
    service.start(upload, UploadStatus.WRITING)
    upload_write.delay(upload.pk)
    upload.refresh_from_db()
    assert upload.status == UploadStatus.DONE, upload.error
    return upload


class TestRegionUpload:
    def test_brazil_book_creates_list_from_br_sheet(
        self, convertio: Product, us: AnchorList
    ) -> None:
        upload = _check(_create(convertio, _book(EN, BR), "br"))
        summary = upload.summary or {}
        assert (summary["region"], summary["new_list"]) == ("Бразилия", True)
        assert summary["sheet"] == "Распределение анкоров BR"
        assert summary["positions_country"] == "BR"
        assert any("безанкорка основного списка" in note for note in summary["notes"])
        _write(upload)
        br = AnchorList.objects.get(product=convertio, country="br")
        found = dict(br.keywords.values_list("keyword", "target_url"))
        assert found == {
            "convert": "https://convertio.co/pt/",
            "conversor de vídeo": "https://convertio.co/pt/video-converter/",
        }
        # «convert» США не тронут — у Бразилии свой.
        assert us.keywords.get(keyword="convert").target_url == "https://convertio.co/"
        assert set(KeywordPosition.objects.values_list("country", flat=True)) == {"BR"}
        # Цели — общие для продукта, безанкорка без пометки — не Бразилии.
        assert PageTypeShare.objects.filter(product=convertio).count() == 1
        assert CountryShare.objects.filter(product=convertio).count() == 3
        assert not br.keywords.filter(keyword="Convertio").exists()

    def test_us_book_takes_en_sheet(self, convertio: Product, us: AnchorList) -> None:
        upload = _check(_create(convertio, _book(EN, BR), "us"))
        summary = upload.summary or {}
        assert (summary["sheet"], summary["new_list"], summary["notes"]) == (
            "Распределение анкоров EN",
            False,
            [],
        )
        _write(upload)
        assert set(us.keywords.values_list("keyword", flat=True)) == {
            "convert",
            "old us key",
            "mp4 to mp3",
            "Convertio",
        }
        assert not AnchorList.objects.filter(country="br").exists()

    def test_two_unmarked_sheets_ask_to_split(self, convertio: Product, us: AnchorList) -> None:
        sheets = (("Анкоры 1", EN[1]), ("Анкоры 2", BR[1]))
        upload = _create(convertio, _book(*sheets, shares=False), "br")
        assert upload.status == UploadStatus.FAILED
        assert "несколько листов анкоров" in (upload.error or "")

    def test_file_of_other_region_is_refused(self, convertio: Product, us: AnchorList) -> None:
        csv = b"Keyword;URL\nconvert;https://convertio.co/\n"
        upload = _create(convertio, csv, "br", name="Линкбилдинг - Распределение анкоров EN.csv")
        assert upload.status == UploadStatus.FAILED
        assert "Похоже, это анкоры региона «США»" in (upload.error or "")

    def test_brazil_file_is_not_loaded_into_usa(self, convertio: Product, us: AnchorList) -> None:
        # Списка Бразилии ещё нет, регион забыли выбрать: пометка «BR» в имени не пускает.
        csv = b"Keyword;URL\nconverter;https://convertio.co/pt/\n"
        name = "Линкбилдинг Convertio.co.xlsx - Распределение анкоров BR.csv"
        upload = _create(convertio, csv, "us", name=name)
        assert upload.status == UploadStatus.FAILED
        assert "Похоже, это анкоры региона «Бразилия»" in (upload.error or "")
        brazil = _check(_create(convertio, csv, "br", name=name))
        assert (brazil.summary or {})["new_total"] == 1

    def test_full_list_disables_only_own_region(self, convertio: Product, us: AnchorList) -> None:
        _write(_check(_create(convertio, _book(BR, shares=False), "br")))
        sync = anchors.Policy.SYNC
        data = _book(("Распределение анкоров BR", BR[1][:1]), shares=False)
        upload = _check(_create(convertio, data, "br", anchor_policy=sync))
        assert (upload.summary or {})["missing"] == ["conversor de vídeo"]
        _write(upload)
        assert us.keywords.get(keyword="old us key").is_active
        assert not Keyword.objects.get(keyword="conversor de vídeo").is_active

    def test_same_file_for_other_region_is_new_upload(
        self, convertio: Product, us: AnchorList
    ) -> None:
        # Файл без пометки региона: тот же файл под другим регионом — другая загрузка.
        data = b"Keyword;URL\nconversor;https://convertio.co/pt/\n"
        first = _write(_check(_create(convertio, data, "br", name="anchors.csv")))
        again = service.create_upload(
            SimpleUploadedFile("anchors.csv", data),
            kind=UploadKind.ANCHORS,
            seller=None,
            product=convertio,
            country="br",
            prices_date=dt.date(2026, 10, 24),
        )
        assert again.duplicate and again.upload == first
        other = _check(_create(convertio, data, "pt", name="anchors.csv"))
        assert other.pk != first.pk
        assert (other.summary or {})["region"] == "Португалия"

    def test_undo_removes_new_list(self, convertio: Product, us: AnchorList) -> None:
        upload = _write(_check(_create(convertio, _book(BR, shares=False), "br")))
        assert AnchorList.objects.filter(country="br").exists()
        undo.undo(upload)
        assert not AnchorList.objects.filter(country="br").exists()
        assert set(Keyword.objects.values_list("keyword", flat=True)) == {"convert", "old us key"}


class TestUploadForm:
    def test_region_picker_with_lists_first(
        self, admin_client: Client, admin_user: User, convertio: Product, us: AnchorList
    ) -> None:
        choose_product(admin_user, convertio)
        url = reverse("admin:sites_upload_add") + "?kind=anchors&anchor_country=br"
        html = admin_client.get(url).content.decode()
        assert 'name="anchor_country"' in html
        lists = html[html.index('label="Уже есть списки анкоров"') :]
        lists = lists[: lists.index("</optgroup>")]
        assert 'value="us"' in lists and 'value="br"' not in lists
        assert 'value="br" selected' in html

    def test_upload_keeps_region(
        self, admin_client: Client, admin_user: User, convertio: Product, us: AnchorList
    ) -> None:
        choose_product(admin_user, convertio)
        admin_client.post(
            reverse("admin:sites_upload_add"),
            {
                "kind": "anchors",
                "product": convertio.pk,
                "anchor_country": "br",
                "anchor_mode": "auto",
                "anchor_policy": "merge",
                "prices_date": "2026-10-24",
                "file": SimpleUploadedFile("br.xlsx", _book(BR, shares=False)),
            },
        )
        upload = Upload.objects.get()
        assert upload.country == "br"
        listed = admin_client.get(reverse("admin:sites_upload_changelist")).content.decode()
        assert "Бразилия" in listed


class TestPlacementsFile:
    def test_links_bind_to_site_region(self, convertio: Product, us: AnchorList) -> None:
        br = AnchorList.objects.create(product=convertio, country="br")
        brand_br = Keyword.objects.create(
            anchor_list=br, keyword="Convertio", target_url="https://convertio.co/pt/"
        )
        brand_us = Keyword.objects.create(
            anchor_list=us, keyword="Convertio", target_url="https://convertio.co/"
        )
        headers = ["Target", "Статус", "Анкор1", "Ссылка1"]
        rows = [
            ["dci.com.br", "Заявка отправлена", "Convertio", "https://convertio.co/pt/"],
            ["example.com", "Заявка отправлена", "Convertio", "https://convertio.co/"],
        ]
        text = "\n".join(";".join(row) for row in [headers, *rows]) + "\n"
        Site.objects.create(domain="example.com", language="en")
        upload = service.create_upload(
            SimpleUploadedFile("placements.csv", text.encode()),
            kind=UploadKind.PLACEMENTS,
            seller=None,
            product=convertio,
            prices_date=dt.date(2026, 10, 24),
        ).upload
        if service.needs_questions(upload):
            assert service.confirm_mapping(upload, upload.mapping or {}, "EUR") == []
        _write(_check(upload))
        brazil = Placement.objects.get(site__domain="dci.com.br")
        usa = Placement.objects.get(site__domain="example.com")
        assert (brazil.anchor_list, brazil.links.get().keyword) == (br, brand_br)
        assert (usa.anchor_list, usa.links.get().keyword) == (us, brand_us)
