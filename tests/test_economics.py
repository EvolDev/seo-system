"""Экономика: бюджет месяца, потрачено, запланировано, остаток и ссылки по месяцам."""

import datetime as dt
import io
from collections.abc import Callable, Iterator
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from django.contrib.auth.models import Permission, User
from django.test import Client, RequestFactory, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.content.domain_settings import MonthlyBudget, set_product_setting
from apps.content.models import DomainSetting
from apps.placements import economics, economics_view
from apps.placements.economics import Period
from apps.placements.models import Placement, PlacementStatus
from apps.sites.models import ExchangeRate, Product, Site, SitePrice

pytestmark = pytest.mark.django_db

OfferFactory = Callable[..., SitePrice]

MSK = ZoneInfo("Europe/Moscow")
TODAY = dt.date(2026, 10, 14)
OCTOBER = Period(dt.date(2026, 10, 1), dt.date(2026, 11, 1))
URL = reverse("admin:economics")


@pytest.fixture(autouse=True)
def moscow() -> Iterator[None]:
    with override_settings(TIME_ZONE="Europe/Moscow"):
        yield


@pytest.fixture
def today(monkeypatch: pytest.MonkeyPatch) -> dt.date:
    monkeypatch.setattr(timezone, "localdate", lambda *args: TODAY)
    return TODAY


@pytest.fixture
def products() -> tuple[Product, Product]:
    convertio = Product.objects.create(name="Convertio", domain="convertio.co")
    clideo = Product.objects.create(name="Clideo", domain="clideo.com")
    # Порог написания нужен, чтобы посчитать, сколько заплатим площадке (ADR-043).
    DomainSetting.objects.create(key="PRICE_REFERENCE", value={"writing_eur": 50})
    return convertio, clideo


def _placement(
    product: Product,
    domain: str,
    status: PlacementStatus,
    *,
    ordered: dt.datetime | None = None,
    published: dt.datetime | None = None,
    paid: int | None = None,
    currency: str = "EUR",
    site: Site | None = None,
) -> Placement:
    return Placement.objects.create(
        site=site or Site.objects.create(domain=domain),
        product=product,
        status=status,
        ordered_at=ordered,
        published_at=published,
        price_paid_cents=paid,
        currency=currency,
    )


@pytest.fixture
def data(products: tuple[Product, Product], offer: OfferFactory) -> tuple[Product, Product]:
    convertio, clideo = products
    ExchangeRate.objects.create(
        currency="USD", rate_date=dt.date(2026, 10, 1), rate=Decimal("1.25")
    )
    # Потрачено: сентябрь и октябрь по дате публикации; доллары — по курсу.
    _placement(
        convertio, "sep.com", PlacementStatus.PLACED,
        published=dt.datetime(2026, 9, 11, 12, tzinfo=MSK), paid=10000,
    )  # fmt: skip
    _placement(
        convertio, "usd.com", PlacementStatus.PLACED,
        published=dt.datetime(2026, 10, 2, 12, tzinfo=MSK), paid=12500, currency="USD",
    )  # fmt: skip
    _placement(
        convertio, "nopaid.com", PlacementStatus.PLACED,
        published=dt.datetime(2026, 10, 3, 12, tzinfo=MSK),
    )  # fmt: skip
    _placement(convertio, "nodate.com", PlacementStatus.PLACED, paid=7000)
    # План: рабочая цена — размещение + анонс + написание до порога (пишет площадка).
    ordered = Site.objects.create(domain="ordered.com")
    offer(ordered, 20000, announce_cents=5000, writing_cents=3000)
    _placement(
        convertio, "", PlacementStatus.ORDERED, site=ordered,
        ordered=dt.datetime(2026, 10, 5, 12, tzinfo=MSK),
    )  # fmt: skip
    # Написание дороже порога — пишем мы, площадке за него не платим.
    in_work = Site.objects.create(domain="inwork.com")
    offer(in_work, 30000, writing_cents=9000)
    _placement(convertio, "", PlacementStatus.IN_WORK, site=in_work)
    # «Заплачено» уже вписано — договорённая сумма главнее рабочей цены.
    _placement(convertio, "agreed.com", PlacementStatus.WRITING, paid=15000)
    # Без рабочей цены — суммы нет; заявка из сентября — в сентябре.
    _placement(
        convertio, "noprice.com", PlacementStatus.ORDERED,
        ordered=dt.datetime(2026, 9, 20, 12, tzinfo=MSK),
    )  # fmt: skip
    # Отказ в план не входит; Clideo — другой продукт.
    _placement(convertio, "rejected.com", PlacementStatus.REJECTED, paid=99999)
    _placement(
        clideo, "clideo.com", PlacementStatus.PLACED,
        published=dt.datetime(2026, 10, 7, 12, tzinfo=MSK), paid=50000,
    )  # fmt: skip
    return convertio, clideo


def test_october(data: tuple[Product, Product]) -> None:
    convertio, _clideo = data
    report = economics.report(OCTOBER, [(convertio.pk, "Convertio")], convertio.pk, TODAY)
    total = report.total
    assert total.budget_cents == 3_000_000
    # usd.com $125 / 1.25 = €100; nopaid.com без суммы.
    assert (total.placed, total.spent_cents, total.priced) == (2, 10000, 1)
    assert total.converted
    # ordered.com €200 + анонс €50 + написание €30; inwork.com €300 без написания;
    # agreed.com — «Заплачено» €150.
    assert (total.planned, total.planned_cents) == (3, 28000 + 30000 + 15000)
    assert total.left_cents == 3_000_000 - 10000 - 73000
    assert total.average_cents == 10000
    statuses = {group.name: (group.count, group.cents) for group in report.by_status}
    assert statuses == {
        "В работе": (1, 30000),
        "Заявка отправлена": (1, 28000),
        "Написание статьи": (1, 15000),
    }
    assert (report.no_paid, report.no_order_date, report.no_published) == (1, 2, 1)
    # noprice.com — в сентябре: заявку отправили тогда.
    assert report.no_price == 0
    domains = [row.domain for row in report.rows]
    assert "rejected.com" not in domains and "clideo.com" not in domains


def test_months_cover_the_work(data: tuple[Product, Product]) -> None:
    convertio, _clideo = data
    report = economics.report(OCTOBER, [(convertio.pk, "Convertio")], convertio.pk, TODAY)
    months = {month.month: month for month in report.months}
    assert list(months) == [dt.date(2026, 9, 1), dt.date(2026, 10, 1)]
    september = months[dt.date(2026, 9, 1)]
    assert (september.placed, september.spent_cents, september.planned) == (1, 10000, 1)
    september_report = economics.report(
        Period(dt.date(2026, 9, 1), dt.date(2026, 10, 1)),
        [(convertio.pk, "Convertio")],
        convertio.pk,
        TODAY,
    )
    assert september_report.no_price == 1


def test_all_products_budget_is_sum(data: tuple[Product, Product]) -> None:
    convertio, clideo = data
    set_product_setting(clideo.pk, "MONTHLY_BUDGET", {"eur": 10000})
    products = [(convertio.pk, "Convertio"), (clideo.pk, "Clideo")]
    report = economics.report(OCTOBER, products, None, TODAY)
    assert report.budgets == [("Convertio", 3_000_000), ("Clideo", 1_000_000)]
    assert report.total.budget_cents == 4_000_000
    assert report.total.placed == 3
    quarter = economics.report(
        Period(dt.date(2026, 10, 1), dt.date(2027, 1, 1)), products, None, TODAY
    )
    assert quarter.total.budget_cents == 3 * 4_000_000


def test_budget_setting_shape() -> None:
    assert MonthlyBudget.parse({"eur": 30000}).eur_cents == 3_000_000
    for wrong in ({"eur": -1}, {"eur": "30000"}, {"sum": 1}, 30000):
        with pytest.raises(ValueError):
            MonthlyBudget.parse(wrong)


class TestPeriod:
    def state(self, **params: str) -> economics_view.State:
        request = RequestFactory().get(URL, params)
        request.user = User(is_superuser=True)
        return economics_view.state_from(request, TODAY)

    def test_default_is_this_month(self, products: tuple[Product, Product]) -> None:
        state = self.state(product=str(products[0].pk))
        assert (state.kind, state.period.start, state.period.end) == (
            "month",
            dt.date(2026, 10, 1),
            dt.date(2026, 11, 1),
        )
        assert state.label == "Октябрь 2026"

    def test_quarter_and_year(self, products: tuple[Product, Product]) -> None:
        quarter = self.state(period="quarter", at="2026-11")
        assert quarter.period == Period(dt.date(2026, 10, 1), dt.date(2027, 1, 1))
        assert quarter.label == "IV квартал 2026"
        year = self.state(period="year", at="2025-05")
        assert year.period == Period(dt.date(2025, 1, 1), dt.date(2026, 1, 1))
        assert year.label == "2025 год"

    def test_range_is_ordered_and_capped(self, products: tuple[Product, Product]) -> None:
        state = self.state(period="range", **{"from": "2026-10", "to": "2025-05"})
        assert state.period == Period(dt.date(2025, 5, 1), dt.date(2026, 11, 1))
        assert state.label == "Май 2025 — октябрь 2026"
        long = self.state(period="range", **{"from": "2020-01", "to": "2026-10"})
        assert len(long.period.months) == economics_view.RANGE_LIMIT

    def test_all_products(self, products: tuple[Product, Product]) -> None:
        state = self.state(product="all")
        assert state.product is None and state.product_label == "Все продукты"


def test_charts(data: tuple[Product, Product]) -> None:
    convertio, _clideo = data
    report = economics.report(OCTOBER, [(convertio.pk, "Convertio")], convertio.pk, TODAY)
    money = economics_view.money_chart(report)
    assert money.budget_y is not None and money.budget_label == "бюджет €30\u202f000"
    october = money.columns[-1]
    assert [segment.key for segment in october.segments] == ["spent", "planned"]
    assert october.label == "€830"
    links = economics_view.links_chart(report)
    assert [point.value for point in links.points] == [1, 2]
    assert links.line.startswith("M")
    assert "Размещено: 2 ссылки" in links.slots[-1].tip


def test_money_texts() -> None:
    assert economics_view.euros(-1266800) == "−€12\u202f668"
    assert economics_view.euros(12500, approx=True) == "≈€125"
    assert economics_view.compact(4266800) == "€42,7k"
    assert economics_view.count_text(21, economics_view.LINKS_FORMS) == "21 ссылка"
    assert economics_view.count_text(12, economics_view.LINKS_FORMS) == "12 ссылок"


class TestScreen:
    def test_page(
        self, admin_client: Client, data: tuple[Product, Product], today: dt.date
    ) -> None:
        convertio, _clideo = data
        response = admin_client.get(URL, {"product": convertio.pk})
        assert response.status_code == 200
        html = response.content.decode()
        assert "Октябрь 2026" in html
        assert "€30\u202f000" in html  # бюджет
        assert "Деньги по месяцам" in html and "Размещено ссылок по месяцам" in html
        assert 'data-econ-filter="nopaid"' in html
        assert "ordered.com" in html and "rejected.com" not in html
        # Без копирайтеров колонки копирайта не повторяют «Потрачено» и «Ср. цену».
        assert "Потрачено + копирайт" not in html and "Колонки копирайта появятся" in html

    def test_menu_and_home_card(self, admin_client: Client, data: tuple[Product, Product]) -> None:
        html = admin_client.get(reverse("admin:index")).content.decode()
        assert f'href="{URL}"' in html
        assert "Экономика" in html

    def test_without_permission(self, client: Client, products: tuple[Product, Product]) -> None:
        user = User.objects.create_user("viewer", password="x", is_staff=True)
        user.user_permissions.add(Permission.objects.get(codename="view_site"))
        client.force_login(user)
        assert client.get(URL).status_code == 403
        assert f'href="{URL}"' not in client.get(reverse("admin:index")).content.decode()

    def test_excel(
        self, admin_client: Client, data: tuple[Product, Product], today: dt.date
    ) -> None:
        convertio, _clideo = data
        response = admin_client.get(URL, {"product": str(convertio.pk), "export": "xlsx"})
        assert response.status_code == 200
        assert "attachment" in response["Content-Disposition"]
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Сводка", "По месяцам", "Размещения"]
        months = book["По месяцам"]
        assert months["B1"].value == "Размещено ссылок"
        assert months["B3"].value == 2  # октябрь
        assert len(months._charts) == 2  # type: ignore[attr-defined]
        rows = book["Размещения"]
        assert {rows.cell(row=r, column=1).value for r in range(2, rows.max_row + 1)} >= {
            "ordered.com",
            "usd.com",
        }

    def test_html_file(
        self, admin_client: Client, data: tuple[Product, Product], today: dt.date
    ) -> None:
        convertio, _clideo = data
        response = admin_client.get(URL, {"product": str(convertio.pk), "export": "html"})
        assert response.status_code == 200
        assert response["Content-Type"].startswith("text/html")
        assert "attachment" in response["Content-Disposition"]
        html = response.content.decode()
        assert html.startswith("<!doctype html>")
        # Стили и скрипт — внутри: файл открывается без системы.
        assert ".seo-econ-tiles" in html and "seoEconomics" in html
        assert "/static/" not in html and "data-panel" not in html
