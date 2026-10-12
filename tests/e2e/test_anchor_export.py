"""Копирование анкоров и выгрузка всей выборки в настоящем браузере."""

from pathlib import Path

import pytest
from django.contrib.auth.models import User
from openpyxl import load_workbook
from playwright.sync_api import Page, expect
from pytest_django.live_server_helper import LiveServer

from apps.keywords.models import Keyword
from apps.sites.models import Product
from apps.workspace.products import choose_product

from .conftest import mark, same_document, soft_load

pytestmark = [pytest.mark.e2e, pytest.mark.django_db(transaction=True, serialized_rollback=True)]


@pytest.fixture
def product(admin_user: User) -> Product:
    product = Product.objects.create(name="Convertio", domain="convertio.co")
    choose_product(admin_user, product)
    Keyword.objects.bulk_create(
        Keyword(
            product=product,
            keyword=f"audio {i:03}",
            target_url="https://convertio.co/wav-mp3/",
            volume=i,
            global_volume=i * 2,
            tool="Audio",
        )
        for i in range(548)
    )
    Keyword.objects.create(
        product=product, keyword="image", target_url="https://convertio.co/png-jpg/", tool="Image"
    )
    return product


def test_copy_all_and_subset(admin_page: Page, live_server: LiveServer, product: Product) -> None:
    page = admin_page
    page.context.grant_permissions(["clipboard-read", "clipboard-write"])
    page.goto(
        f"{live_server.url}/admin/keywords/keywordcoverage/?product={product.pk}&tool=Audio&per_page=500"
    )
    expect(page.locator("#result_list tbody tr")).to_have_count(500)
    mark(page)
    page.locator("#action-toggle").check()
    page.locator(".actions .question a").click()
    expect(page.locator('input[name="select_across"]')).to_have_value("1")
    button = page.get_by_role("button", name="Скопировать анкоры", exact=True)
    button.click()
    expect(page.get_by_role("button", name="Скопировано", exact=True)).to_be_visible()
    text = page.evaluate("navigator.clipboard.readText()")
    assert text.splitlines() == [f"audio {i:03}" for i in range(548)]
    assert same_document(page)
    page.locator(".actions").screenshot(
        path="/app/.anchor-export-previews/copied-buttons.png", animations="disabled"
    )
    page.locator(".actions .clear a").click()
    page.locator('input[name="_selected_action"]').nth(3).check()
    expect(button).to_be_visible()
    button.click()
    expect(page.get_by_role("button", name="Скопировано", exact=True)).to_be_visible()
    assert page.evaluate("navigator.clipboard.readText()") == "audio 003"


def test_excel_filters_after_soft_navigation(
    admin_page: Page, live_server: LiveServer, product: Product, tmp_path: Path
) -> None:
    page = admin_page
    page.goto(
        f"{live_server.url}/admin/keywords/keywordcoverage/?product={product.pk}&tool=Audio&per_page=500"
    )
    mark(page)
    # Export ignores individual selection: all filtered rows, as the button promises.
    page.locator('input[name="_selected_action"]').first.check()
    with page.expect_download() as transfer:
        page.get_by_role("link", name="Экспорт в Excel", exact=True).click()
    file = tmp_path / "anchors.xlsx"
    transfer.value.save_as(file)
    sheet = load_workbook(file)["Анкоры"]
    assert sheet.max_row == 549
    assert [cell.value for cell in sheet[1]] == ["Keyword", "URL", "Volume", "Global Volume"]
    assert sheet["B2"].value == "https://convertio.co/wav-mp3/"
    assert same_document(page)
    page.locator("#searchbar").fill("audio 547")
    with soft_load(page):
        page.get_by_role("button", name="Найти", exact=True).click()
    expect(page.locator("#result_list tbody tr")).to_have_count(1)
    with page.expect_download() as filtered:
        page.get_by_role("link", name="Экспорт в Excel", exact=True).click()
    filtered.value.save_as(file)
    sheet = load_workbook(file)["Анкоры"]
    assert list(sheet.values) == [
        ("Keyword", "URL", "Volume", "Global Volume"),
        ("audio 547", "https://convertio.co/wav-mp3/", 547, 1094),
    ]
    assert same_document(page)
