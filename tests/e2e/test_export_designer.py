"""Реальный браузер: конструктор, сохранение, перетаскивание, импорт и оба файла."""

import csv
import datetime as dt
import io
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from playwright.sync_api import Page, expect
from pytest_django.live_server_helper import LiveServer

from apps.placements.models import Placement
from apps.sites.models import ExchangeRate, Product, Site

pytestmark = [pytest.mark.e2e, pytest.mark.django_db(transaction=True, serialized_rollback=True)]
URL = "/admin/placements/placement/export-designer/"


@pytest.fixture
def placements() -> None:
    product = Product.objects.create(name="Convertio", domain="convertio.co")
    for number in range(7):
        Placement.objects.create(
            site=Site.objects.create(domain=f"site{number}.com"),
            product=product,
            price_paid_cents=10000 + number,
        )


def test_constructor_profile_drag_and_files(
    admin_page: Page, live_server: LiveServer, placements: None, tmp_path: Path
) -> None:
    page = admin_page
    page.goto(live_server.url + URL)
    expect(page.locator("#server-status")).to_contain_text("Готово: 7 строк")
    expect(page.locator("#columns tr")).to_have_count(35)
    assert page.locator("#unlock-drag").count() == 0
    page.locator('[data-preset="anchors"]').click()
    expect(page.locator("#server-status")).to_contain_text("5 колонок")
    page.locator("#columns tr").first.locator('[data-action="header"]').fill("Донор")
    page.locator("#columns tr").first.locator('[data-action="header"]').press("Tab")
    page.locator("#presetname").fill("Мой отчёт")
    page.locator("#saveform button").click()
    expect(page.locator("#savedlist button")).to_contain_text("Мой отчёт")
    page.reload()
    page.locator("#savedlist button").click()
    expect(page.locator("#columns tr").first.locator('[data-action="header"]')).to_have_value(
        "Донор"
    )
    page.locator("#expand-columns").click()
    expect(page.locator("#columns-dialog")).to_be_visible()
    before = page.locator("#columns tr").evaluate_all("rs=>rs.map(r=>r.dataset.id)")
    handle = page.locator("#columns tr").first.locator(".drag-handle").bounding_box()
    target = page.locator("#columns tr").nth(3).bounding_box()
    assert handle is not None and target is not None
    page.mouse.move(handle["x"] + 8, handle["y"] + 8)
    page.mouse.down()
    page.mouse.move(handle["x"] + 8, target["y"] + target["height"] * 0.8, steps=18)
    page.mouse.up()
    expect(page.locator("#columns tr").first).not_to_have_attribute("data-id", before[0])
    page.locator("#expand-columns").click()
    expect(page.locator("#expand-columns")).to_have_attribute("aria-expanded", "false")
    page.locator("#tab-preview").click()
    expect(page.locator("#preview tbody tr")).to_have_count(5)
    with (
        page.expect_response(lambda r: "attachment" in r.headers.get("content-disposition", "")),
        page.expect_download() as info,
    ):
        page.locator("#download-xlsx").click()
    path = tmp_path / "export.xlsx"
    info.value.save_as(path)
    sheet = load_workbook(path).active
    assert sheet is not None and sheet.max_row == 8 and sheet.max_column == 5
    assert "Донор" in [cell.value for cell in sheet[1]]
    with (
        page.expect_response(lambda r: "attachment" in r.headers.get("content-disposition", "")),
        page.expect_download() as info,
    ):
        page.locator("#download").click()
    path = tmp_path / "export.csv"
    info.value.save_as(path)
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig")), delimiter=";"))
    assert len(rows) == 8 and len(rows[0]) == 5


def test_import_order_and_resolve_missing(
    admin_page: Page, live_server: LiveServer, placements: None, tmp_path: Path
) -> None:
    page = admin_page
    page.goto(live_server.url + URL)
    expect(page.locator("#server-status")).to_contain_text("Готово")
    sample = tmp_path / "sample.csv"
    sample.write_text(",Месяц\nsite1.com,Oct-26\nmissing.com,Oct-26\nsite0.com,Sep-26\n")
    page.locator("#samplefile").set_input_files(sample)
    page.locator("#importform button").click()
    expect(page.locator("#server-status")).to_contain_text("Нужно разобрать 1 строк")
    page.locator("#tab-issues").click()
    page.locator('[data-line="3"]').select_option("keep")
    expect(page.locator("#server-status")).to_contain_text("Готово: 3 строк")
    page.locator("#tab-preview").click()
    expect(page.locator("#preview tbody tr").first.locator("td").nth(1)).to_have_text("site1.com")
    with (
        page.expect_response(lambda r: "attachment" in r.headers.get("content-disposition", "")),
        page.expect_download() as info,
    ):
        page.locator("#download-xlsx").click()
    path = tmp_path / "filled.xlsx"
    info.value.save_as(path)
    sheet = load_workbook(path).active
    assert sheet is not None
    assert [sheet.cell(i, 1).value for i in (2, 3, 4)] == ["site1.com", "missing.com", "site0.com"]
    assert sheet["B2"].value == "Oct-26"


def test_selected_rows_from_list(
    admin_page: Page, live_server: LiveServer, placements: None
) -> None:
    page = admin_page
    page.goto(live_server.url + "/admin/placements/placement/")
    page.locator("#result_list input.action-select").nth(0).check()
    page.locator("#result_list input.action-select").nth(2).check()
    page.locator("[data-designer-open] button").click()
    expect(page.locator("#server-status")).to_contain_text("Готово: 2 строк")


def test_constructor_admin_palettes_and_soft_navigation(
    admin_page: Page, live_server: LiveServer, placements: None
) -> None:
    page = admin_page
    page.goto(live_server.url + URL)
    expect(page.locator("#server-status")).to_contain_text("Готово: 7 строк")
    expect(page.locator("#header")).to_be_visible()
    expect(page.locator("#nav-sidebar")).to_be_visible()
    colors = set()
    for palette in ("apple-light", "google-dark", "emerald", "ahrefs"):
        page.locator("#palette-switch summary").click()
        page.locator(f'#palette-switch button[data-palette="{palette}"]').click()
        expect(page.locator("html")).to_have_attribute("data-palette", palette)
        page.mouse.move(500, 120)
        page.wait_for_timeout(200)  # Штатный переход цветов кнопок длится 150 мс.
        colors.add(
            page.locator("#download-xlsx").evaluate("e=>getComputedStyle(e).backgroundColor")
        )
        body = page.locator("body").evaluate("e=>getComputedStyle(e).backgroundColor")
        assert page.locator("#search").evaluate("e=>getComputedStyle(e).backgroundColor") == body
        page.locator("#expand-columns").click()
        expect(page.locator("#columns-dialog")).to_be_visible()
        assert (
            page.locator("#columns-dialog").evaluate("e=>getComputedStyle(e).backgroundColor")
            == body
        )
        page.locator("#expand-columns").click()
    assert len(colors) == 4
    page.evaluate("window.__exportSameDocument = true")
    for _ in range(2):
        page.locator(".designer-back a").click()
        expect(page.locator("#result_list")).to_be_visible()
        page.evaluate("url=>window.seoNav.visit(url)", live_server.url + URL)
        expect(page.locator("#server-status")).to_contain_text("Готово: 7 строк")
        page.locator("#expand-columns").click()
        expect(page.locator("#columns-dialog")).to_be_visible()
        page.locator("#expand-columns").click()
        assert page.evaluate("window.__exportSameDocument === true")


def test_currency_choice_converts_and_is_remembered(
    admin_page: Page, live_server: LiveServer
) -> None:
    """Два выбора валюты: пересчёт в предпросмотре, выбор — после перезагрузки."""
    product = Product.objects.create(name="Convertio", domain="convertio.co")
    Placement.objects.create(
        site=Site.objects.create(domain="dollar.com"),
        product=product,
        price_paid_cents=10000,
        currency="USD",
    )
    ExchangeRate.objects.create(
        currency="USD", rate_date=dt.date(2026, 10, 16), rate=Decimal("1.25")
    )
    page = admin_page
    page.goto(live_server.url + URL)
    expect(page.locator("#server-status")).to_contain_text("Готово: 1 строк")
    expect(page.locator("#currency-from")).to_have_value("")
    expect(page.locator("#currency-to")).to_have_value("")
    paid = page.locator(
        "#columns tr", has=page.locator('[data-action="header"][value="Итог цена"]')
    )
    expect(paid.locator(".sample-value")).to_have_text("100.00")

    page.locator("#currency-to").select_option("EUR")
    expect(paid.locator(".sample-value")).to_have_text("80.00")
    expect(page.locator("#currencyhelp")).to_contain_text("в EUR")
    page.wait_for_timeout(1000)
    page.reload()
    expect(page.locator("#currency-to")).to_have_value("EUR")
    expect(paid.locator(".sample-value")).to_have_text("80.00")
