"""Выгрузка в Excel и CSV в браузере (E1-06, ADR-053): кнопки над списком
скачивают файл без ухода со страницы и берут фильтры, выбранные без перезагрузки."""

import csv
import datetime as dt
import io
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from playwright.sync_api import Locator, Page, expect
from pytest_django.live_server_helper import LiveServer

from apps.placements.models import Placement, PlacementStatus
from apps.sites.models import Product, Site, SiteList, SiteListItem, SiteMetric

from .conftest import mark, same_document, soft_load

pytestmark = [pytest.mark.e2e, pytest.mark.django_db(transaction=True, serialized_rollback=True)]

SITES = "/admin/sites/productsitelatest/"
PLACEMENTS = "/admin/placements/placement/"


@pytest.fixture
def sites() -> None:
    """150 площадок в рабочем списке — две страницы; у двадцати опубликовано в сентябре."""
    convertio = Product.objects.create(name="Convertio", domain="convertio.co")
    site_list = SiteList.objects.create(name="Октябрь")
    for number in range(150):
        site = Site.objects.create(domain=f"site{number:03}.com")
        SiteMetric.objects.create(site=site, dr=number % 100)
        SiteListItem.objects.create(site_list=site_list, site=site)
        if number < 20:
            Placement.objects.create(
                site=site,
                product=convertio,
                status=PlacementStatus.PUBLISHED,
                published_at=dt.datetime(2026, 9, 1 + number, 12, tzinfo=dt.UTC),
                price_paid_cents=10000 + number,
            )


def _window(page: Page) -> Locator:
    """Открыть окно выгрузки над списком."""
    page.locator(".object-tools a", has_text="Выгрузить").click()
    window = page.locator("[data-export-form]")
    expect(window).to_be_visible()
    return window


def _columns(window: Locator, *titles: str) -> None:
    """Оставить отмеченными только эти колонки."""
    everything = window.locator("[data-export-all]")
    everything.check()
    everything.uncheck()
    for title in titles:
        window.locator("label", has_text=re.compile(f"^ ?{re.escape(title)}$")).locator(
            "input"
        ).check()


def test_sites_export_window(
    admin_page: Page, live_server: LiveServer, sites: None, tmp_path: Path
) -> None:
    page = admin_page
    page.goto(f"{live_server.url}{SITES}")
    mark(page)
    # Фильтр без перезагрузки — окно уже с ним.
    with soft_load(page):
        page.locator("#id_dr__range__gte").fill("50")
        page.locator("#id_dr__range__gte").press("Enter")
    assert "dr__range__gte=50" in page.url

    window = _window(page)
    expect(window.locator("[data-export-rows]")).to_have_text("Строки: все отобранные — 50")
    # Ни одной колонки — не выгружается, окно говорит почему.
    window.locator("[data-export-all]").uncheck()
    window.locator("button", has_text="Excel").click()
    expect(window.locator("[data-export-error]")).to_be_visible()

    _columns(window, "Домен", "DR")
    expect(window.locator("[data-export-all]")).not_to_be_checked()
    with page.expect_download() as download_info:
        window.locator("button", has_text="Excel").click()
    download = download_info.value
    assert download.suggested_filename.startswith("Площадки Convertio ")
    assert download.suggested_filename.endswith(".xlsx")
    path = tmp_path / "sites.xlsx"
    download.save_as(path)
    rows = list(load_workbook(path)["Площадки"].iter_rows(values_only=True))
    # DR от 50 — 50 площадок из 150, все страницы; только выбранные колонки.
    assert rows[0] == ("Домен", "DR")
    assert len(rows) == 1 + 50
    expect(window).to_be_hidden()

    # Отмеченные строки — только они; колонки — как запомнились.
    page.locator("#result_list input.action-select").nth(0).check()
    page.locator("#result_list input.action-select").nth(2).check()
    window = _window(page)
    expect(window.locator("[data-export-rows]")).to_have_text("Строки: отмеченные — 2")
    with page.expect_download() as download_info:
        window.locator("button", has_text="CSV").click()
    content = Path(download_info.value.path()).read_bytes()
    assert content.startswith(b"\xef\xbb\xbf")
    table = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert table[0] == ["Домен", "DR"]
    assert len(table) == 1 + 2

    # После перехода окно приходит с запомненными галочками; Esc и щелчок мимо закрывают.
    with soft_load(page):
        page.locator("#nav-sidebar a", has_text="Площадки").click()
    window = _window(page)
    expect(window.locator("input[name=columns][value=dr]")).to_be_checked()
    expect(window.locator("input[name=columns][value=status]")).not_to_be_checked()
    page.keyboard.press("Escape")
    expect(window).to_be_hidden()
    window = _window(page)
    page.locator("#result_list tbody td.field-dr_cell").nth(20).click()
    expect(window).to_be_hidden()
    # Страница та же — без ухода и перезагрузки.
    assert same_document(page)


def test_month_report_from_placements(
    admin_page: Page, live_server: LiveServer, sites: None, tmp_path: Path
) -> None:
    page = admin_page
    page.goto(f"{live_server.url}{PLACEMENTS}")
    mark(page)
    with soft_load(page):
        # «Опубликовано» есть и среди статусов — фильтр ищем по заголовку.
        title = page.locator("h3", has_text=re.compile("^Опубликовано$"))
        dropdown = page.locator(".list-filter-dropdown", has=title)
        dropdown.locator("select").select_option(label="сентябрь 2026")
    assert "month=2026-09" in page.url

    window = _window(page)
    with page.expect_download() as download_info:
        window.locator("button", has_text="Excel").click()
    download = download_info.value
    assert download.suggested_filename == "Размещения сентябрь 2026.xlsx"
    path = tmp_path / "month.xlsx"
    download.save_as(path)
    book = load_workbook(path)
    assert book.sheetnames == ["Размещения", "Итого"]
    assert len(list(book["Размещения"].iter_rows())) == 1 + 20
    total = [cell.value for cell in book["Итого"][2]]
    assert total[:3] == ["Всего", None, 20]
    assert same_document(page)
