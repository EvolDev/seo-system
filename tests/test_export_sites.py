"""Выгрузка «Площадок» в Excel и CSV (E1-06, ADR-053)."""

import csv
import datetime as dt
import io
from collections.abc import Callable
from decimal import Decimal
from typing import Any

import pytest
from django.contrib.auth.models import User
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from apps.sites import export as site_export
from apps.sites.models import (
    ExchangeRate,
    MetricSource,
    Product,
    ProductSite,
    ProductSiteLatest,
    Seller,
    Site,
    SiteCountryMetric,
    SiteList,
    SiteListItem,
    SiteMetric,
    SitePrice,
    SiteStatus,
    ensure_product_sites,
)
from apps.workspace.products import choose_product

pytestmark = pytest.mark.django_db

OfferFactory = Callable[..., SitePrice]
LIST_URL = reverse("admin:sites_productsitelatest_changelist")


def _url(fmt: str = "xlsx") -> str:
    return reverse("admin:sites_productsitelatest_export", args=[fmt])


@pytest.fixture
def base(offer: OfferFactory) -> dict[str, Any]:
    """Convertio и Clideo; 150 площадок (больше страницы); у example.com — два
    продавца, у bad.com — отказ; в свежем списке — только пять площадок."""
    convertio = Product.objects.create(name="Convertio", domain="convertio.co")
    Product.objects.create(name="Clideo", domain="clideo.com")
    sites = Site.objects.bulk_create(Site(domain=f"site{n:03}.com") for n in range(148))
    SiteMetric.objects.bulk_create(SiteMetric(site=s, dr=n % 90) for n, s in enumerate(sites))

    example = Site.objects.create(domain="example.com", language="en", link_type="dofollow")
    SiteMetric.objects.create(
        site=example,
        dr=95,
        organic_traffic=120000,
        top_geo="us",
        top_geo_traffic=60000,
        source=MetricSource.AHREFS_BATCH,
    )
    linkhub = Seller.objects.create(name="LinkHub Media", currency="USD")
    ExchangeRate.objects.create(
        currency="USD", rate_date=dt.date(2026, 9, 30), rate=Decimal("1.1355")
    )
    offer(example, 20000, writing_cents=1000, announce_cents=0)
    offer(example, 17000, seller=linkhub, currency="USD", working=False)

    bad = Site.objects.create(domain="bad.com")
    SiteMetric.objects.create(site=bad, dr=94)
    ensure_product_sites()
    ProductSite.objects.filter(product=convertio, site=bad).update(
        status=SiteStatus.DISCARDED, comment="Nofollow, отбрасываем"
    )

    older = SiteList.objects.create(name="Сентябрь 2026")
    SiteListItem.objects.bulk_create(SiteListItem(site_list=older, site=s) for s in sites)
    fresh = SiteList.objects.create(name="Collaborator · 02.10.2026")
    SiteListItem.objects.bulk_create(SiteListItem(site_list=fresh, site=s) for s in sites[:5])
    return {"example": example, "bad": bad, "fresh": fresh}


def _xlsx(client: Client, **params: str) -> list[list[Any]]:
    response = client.get(_url(), params)
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["Площадки"]
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def _by_domain(rows: list[list[Any]]) -> dict[str, dict[str, Any]]:
    header = rows[0]
    return {row[0]: dict(zip(header, row, strict=True)) for row in rows[1:]}


def test_all_pages_and_screen_order(admin_client: Client, base: dict[str, Any]) -> None:
    """Все отобранные строки, а не сотня первой страницы; порядок — как на экране."""
    rows = _xlsx(admin_client, list="all")
    assert len(rows) == 1 + 150
    domains = [row[0] for row in rows[1:]]
    assert domains[:2] == ["example.com", "bad.com"]  # DR 95, 94
    assert len(set(domains)) == 150


def test_working_price_and_cheaper_offer(admin_client: Client, base: dict[str, Any]) -> None:
    """Рабочая цена с продавцом и услугой и лучшее предложение дешевле неё."""
    row = _by_domain(_xlsx(admin_client, list="all"))["example.com"]
    assert row["Статус"] == "Новая"
    assert row["DR"] == 95
    assert row["Трафик"] == 120000
    assert row["Топ регион"] == "US"
    assert row["Трафик топ региона"] == 60000
    assert row["Цена"] == 200
    assert row["Валюта"] == "EUR"
    assert row["Цена, EUR"] == 200
    assert row["Услуга"] == "публикация"
    assert row["Продавец"] == "Collaborator"
    assert row["Анонс"] == 0
    assert row["Написание"] == 10
    assert row["Дешевле у продавца"] == "LinkHub Media"
    assert row["Дешевле: цена"] == 170
    assert row["Дешевле: валюта"] == "USD"
    # $170 по курсу 1.1355 — €149.71.
    assert row["Дешевле, EUR"] == 149.71
    assert row["Тип ссылки"] == "dofollow"
    assert row["Язык"] == "en"
    # Замер Ahrefs — доверенный: «со слов продавца» пусто.
    assert row["Метрики со слов продавца"] is None


def test_status_filter_and_reason(admin_client: Client, base: dict[str, Any]) -> None:
    rows = _xlsx(admin_client, list="all", status__exact=SiteStatus.DISCARDED)
    assert len(rows) == 2
    row = _by_domain(rows)["bad.com"]
    assert row["Статус"] == "Отбрасываю"
    assert row["Комментарий"] == "Nofollow, отбрасываем"


def test_default_list_like_screen(admin_client: Client, base: dict[str, Any]) -> None:
    """Без выбора списка — все площадки, как на экране (E1-19)."""
    assert len(_xlsx(admin_client)) == len(_xlsx(admin_client, list="all"))


def test_search_by_pasted_address(admin_client: Client, base: dict[str, Any]) -> None:
    rows = _xlsx(admin_client, list="all", q="https://www.example.com/blog/post")
    assert [row[0] for row in rows[1:]] == ["example.com"]


def test_region_columns(admin_client: Client, base: dict[str, Any]) -> None:
    SiteCountryMetric.objects.create(
        site=base["example"], country="us", organic_traffic=53814, total_keywords=13175
    )
    rows = _xlsx(admin_client, list="all", region="us")
    header = rows[0]
    at = header.index("Трафик топ региона")
    assert header[at + 1 : at + 4] == ["Трафик US", "Ключи US", "Замер US"]
    row = _by_domain(rows)["example.com"]
    assert (row["Трафик US"], row["Ключи US"]) == (53814, 13175)
    assert isinstance(row["Замер US"], dt.datetime)


def test_other_product_and_file_name(
    admin_client: Client, admin_user: User, base: dict[str, Any]
) -> None:
    # Продукт выгрузки — рабочий, из шапки: фильтр продукта его не меняет (ADR-063).
    clideo = Product.objects.get(name="Clideo")
    choose_product(admin_user, clideo)
    response = admin_client.get(_url(), {"list": "all"})
    assert (
        "%D0%9F%D0%BB%D0%BE%D1%89%D0%B0%D0%B4%D0%BA%D0%B8%20Clideo"
        in (response["Content-Disposition"])
    )  # «Площадки Clideo …»
    rows = _by_domain(
        [
            [cell.value for cell in row]
            for row in load_workbook(io.BytesIO(response.content))["Площадки"].iter_rows()
        ]
    )
    # Отказ — у Convertio; у Clideo площадка новая.
    assert rows["bad.com"]["Статус"] == "Новая"


def test_csv_same_rows(admin_client: Client, base: dict[str, Any]) -> None:
    response = admin_client.get(_url("csv"), {"list": "all"})
    assert response["Content-Type"] == "text/csv; charset=utf-8"
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert len(rows) == 1 + 150
    row = dict(zip(rows[0], rows[1], strict=True))
    assert row["Домен"] == "example.com"
    assert row["Цена"] == "200,00"
    assert row["Дешевле, EUR"] == "149,71"


def test_queries_do_not_grow_with_rows(admin_client: Client, base: dict[str, Any]) -> None:
    # Первый выбор заводит строку настроек (память фильтров, E1-19) — он вне замера.
    admin_client.get(_url(), {"list": "all"})
    with CaptureQueriesContext(connection) as small:
        admin_client.get(_url(), {"list": str(base["fresh"].pk)})
    with CaptureQueriesContext(connection) as big:
        admin_client.get(_url(), {"list": "all"})
    assert len(big) == len(small)


def test_window_carries_current_filters(admin_client: Client, base: dict[str, Any]) -> None:
    """Окно выгрузки: адреса с фильтрами списка, галочки по разделам, сколько строк."""
    page = admin_client.get(LIST_URL, {"list": "all", "status__exact": "new"}).content.decode()
    assert f'href="{_url()}?list=all&amp;status__exact=new"' in page
    assert f'data-csv="{_url("csv")}?list=all&amp;status__exact=new"' in page
    assert "Выгрузить ▾" in page
    assert "Строки: все отобранные — 149" in page
    for group in ("Площадка", "Метрики", "Рабочая цена", "Другие предложения"):
        assert f"<legend>{group}</legend>" in page
    assert 'name="columns" value="dr" checked' in page
    # Колонки региона — только когда регион выбран.
    assert 'value="region"' not in page
    with_region = admin_client.get(LIST_URL, {"list": "all", "region": "us"}).content.decode()
    assert "Регион US: трафик, ключи, замер" in with_region


def _post(
    client: Client, columns: list[str], fmt: str = "xlsx", query: str = "list=all", **data: Any
) -> Any:
    """Отправка окна выгрузки: фильтры списка — в адресе, галочки и строки — в форме."""
    return client.post(f"{_url(fmt)}?{query}", {"columns": columns, **data})


def test_chosen_columns_in_file_order(admin_client: Client, base: dict[str, Any]) -> None:
    """Выбранные колонки — в порядке экрана, а не отметки."""
    response = _post(admin_client, ["placement_eur_cents", "domain", "price_seller"])
    rows = [
        [cell.value for cell in row]
        for row in load_workbook(io.BytesIO(response.content))["Площадки"].iter_rows()
    ]
    assert rows[0] == ["Домен", "Цена, EUR", "Продавец"]
    assert ["example.com", 200, "Collaborator"] in rows
    assert len(rows) == 1 + 150


def test_unchecked_columns_remembered_in_session(
    admin_client: Client, base: dict[str, Any]
) -> None:
    keys = [choice.key for choice in site_export.choices(None)]
    _post(admin_client, [key for key in keys if key not in ("dr", "last_note")])
    page = admin_client.get(LIST_URL, {"list": "all"}).content.decode()
    assert 'name="columns" value="dr">' in page  # без checked
    assert 'name="columns" value="domain" checked' in page
    assert "data-export-all>" in page  # «Все колонки» не отмечена
    # Ссылка без окна (GET) — тоже без снятых колонок.
    header = _xlsx(admin_client, list="all")[0]
    assert "DR" not in header and "Последняя заметка" not in header
    assert "Домен" in header
    # Снова все колонки — галочки вернулись.
    _post(admin_client, keys)
    page = admin_client.get(LIST_URL, {"list": "all"}).content.decode()
    assert 'name="columns" value="dr" checked' in page


def test_region_choice_kept_when_exported_without_region(
    admin_client: Client, base: dict[str, Any]
) -> None:
    """Галочка, которой сейчас нет в окне (региона без региона), остаётся как была."""
    with_region = [choice.key for choice in site_export.choices("us")]
    without = [choice.key for choice in site_export.choices(None)]
    _post(admin_client, [key for key in with_region if key != "region"], query="list=all&region=us")
    _post(admin_client, [key for key in without if key != "dr"])
    page = admin_client.get(LIST_URL, {"list": "all", "region": "us"}).content.decode()
    assert 'name="columns" value="region">' in page  # снята при выгрузке с регионом
    assert 'name="columns" value="dr">' in page
    assert 'name="columns" value="domain" checked' in page


def test_only_picked_rows(admin_client: Client, base: dict[str, Any]) -> None:
    """Отмечены строки галочками — в файле только они, в порядке списка."""
    rows = {
        row.domain: row.pk for row in ProductSiteLatest.objects.filter(product__name="Convertio")
    }
    ids = [str(rows["bad.com"]), str(rows["example.com"])]
    response = _post(admin_client, ["domain", "dr"], ids=ids)
    sheet = load_workbook(io.BytesIO(response.content))["Площадки"]
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [
        ["Домен", "DR"],
        ["example.com", 95],
        ["bad.com", 94],
    ]
    assert "%D0%B2%D1%8B%D0%B1%D1%80%D0%B0%D0%BD%D0%BE%202" in response["Content-Disposition"]
    # «Выбраны все N» — все отобранные, номера не важны.
    across = _post(admin_client, ["domain"], fmt="csv", ids=ids, select_across="1")
    assert len(across.content.decode("utf-8-sig").splitlines()) == 1 + 150


def test_no_columns_is_refused(admin_client: Client, base: dict[str, Any]) -> None:
    assert _post(admin_client, []).status_code == 400
    assert _post(admin_client, ["nonsense"]).status_code == 400


def test_unknown_format_and_access(client: Client, admin_client: Client) -> None:
    Product.objects.create(name="Convertio", domain="convertio.co")
    assert admin_client.get(_url("pdf")).status_code == 404
    # Без входа — на страницу входа.
    assert client.get(_url()).status_code == 302
    User.objects.create_user("viewer", password="x", is_staff=True)
    client.login(username="viewer", password="x")
    assert client.get(_url()).status_code == 403
