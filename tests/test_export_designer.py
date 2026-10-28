"""Контракт конструктора: порядок, файлы, неопределённые строки и личные шаблоны."""

import csv
import datetime as dt
import io
import json
from decimal import Decimal
from typing import Any

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.placements import export_designer as design
from apps.placements.models import Placement, PlacementLink
from apps.sites.models import ExchangeRate, Product, Site
from apps.sites.rates import convert_cents
from apps.workspace.models import ExportProfile

pytestmark = pytest.mark.django_db


@pytest.fixture
def placed() -> Placement:
    product = Product.objects.create(name="Convertio", domain="convertio.co")
    site = Site.objects.create(domain="example.com")
    return Placement.objects.create(
        site=site,
        product=product,
        price_paid_cents=12345,
        currency="EUR",
        ordered_at=dt.datetime(2026, 10, 1, tzinfo=dt.UTC),
    )


def post(client: Client, action: str, **payload: Any) -> Any:
    return client.post(
        reverse("admin:placements_placement_export_designer"),
        {"action": action, "payload": json.dumps(payload)},
    )


def config(*sources: str) -> dict[str, Any]:
    value = design.initial(design.fields())
    if sources:
        value["columns"] = [
            {
                "id": str(i),
                "index": None,
                "header": source,
                "source": source,
                "enabled": True,
                "constant": "",
            }
            for i, source in enumerate(sources)
        ]
    return value


def test_default_and_typed_excel(admin_client: Client, placed: Placement) -> None:
    value = config()
    preview = post(admin_client, "preview", config=value).json()
    # Набор листа и «Region» после «Region Traff» (ADR-076).
    assert len(preview["headers"]) == 35
    assert preview["headers"][0] == "" and preview["headers"][6] == "Дата отправки"
    assert preview["rows"][0][0] == "example.com"
    excel = post(admin_client, "download", config=value, digest=preview["digest"], format="xlsx")
    book = load_workbook(io.BytesIO(excel.content))
    assert book.sheetnames == ["Размещения"]
    assert book.active is not None
    assert book.active["G2"].value == dt.datetime(2026, 10, 1)
    assert book.active["AA2"].value == 123.45
    assert excel["Content-Type"].startswith("application/vnd.openxmlformats")


def test_default_money_like_google_sheet(admin_client: Client, placed: Placement) -> None:
    # Рабочая таблица в формате Financial: точка, две цифры копеек, без знака валюты.
    value = config("paid", "target")
    placed.price_paid_cents = 750
    placed.save()
    preview = post(admin_client, "preview", config=value).json()
    assert preview["rows"] == [["7.50", "example.com"]]
    copied = post(admin_client, "copy", config=value, digest=preview["digest"], start=0, stop=1)
    assert copied.json()["rows"] == [["7.50"]]
    response = post(admin_client, "download", config=value, digest=preview["digest"], format="csv")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[1] == ["7.50", "example.com"]


def test_renamed_reordered_fields_csv_and_formula_safety(
    admin_client: Client, placed: Placement
) -> None:
    value = config("paid", "target", "comment")
    value.update(delimiter=",", decimal=".")
    value["columns"][0]["header"] = "=1+1"
    placed.comment = '=HYPERLINK("bad")'
    placed.save()
    preview = post(admin_client, "preview", config=value).json()
    response = post(admin_client, "download", config=value, digest=preview["digest"], format="csv")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == ["'=1+1", "target", "comment"]
    assert rows[1] == ["123.45", "example.com", '\'=HYPERLINK("bad")']
    response = post(admin_client, "download", config=value, digest=preview["digest"], format="xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None and sheet["A1"].data_type == "s"
    assert sheet["C2"].data_type == "s"


def test_link_slots_do_not_shift(admin_client: Client, placed: Placement) -> None:
    PlacementLink.objects.create(
        placement=placed, link_index=2, anchor="second", target_url="https://convertio.co/"
    )
    preview = post(
        admin_client, "preview", config=config("links.1.anchor", "links.2.anchor")
    ).json()
    assert preview["rows"] == [["", "second"]]


def test_changed_data_requires_new_preview(admin_client: Client, placed: Placement) -> None:
    value = config("paid")
    digest = post(admin_client, "preview", config=value).json()["digest"]
    placed.price_paid_cents = 450
    placed.save()
    response = post(admin_client, "download", config=value, digest=digest, format="xlsx")
    assert response.status_code == 409 and "изменились" in response.json()["error"]


def test_upload_preserves_blank_headers_rows_and_duplicate_domains(
    admin_client: Client, placed: Placement
) -> None:
    raw = b",Month\nexample.com,Oct-26\n,\nexample.com,Sep-26\nmissing.com,Oct-26\n"
    response = admin_client.post(
        reverse("admin:placements_placement_export_designer"),
        {"action": "import", "file": SimpleUploadedFile("sample.csv", raw)},
    )
    imported = response.json()
    value = imported["config"]
    value["columns"][1]["source"] = "original"
    preview = post(admin_client, "preview", config=value, token=imported["token"]).json()
    assert len(preview["issues"]) == 3
    assert preview["count"] == 4 and preview["headers"][0] == ""
    blocked = post(
        admin_client,
        "download",
        config=value,
        token=imported["token"],
        digest=preview["digest"],
        format="xlsx",
    )
    assert blocked.status_code == 409
    decisions = {
        "2": {"action": "match", "id": placed.pk},
        "4": {"action": "keep"},
        "5": {"action": "keep"},
    }
    preview = post(
        admin_client, "preview", config=value, token=imported["token"], decisions=decisions
    ).json()
    assert not preview["issues"] and preview["rows"][1] == ["", ""]
    response = post(
        admin_client,
        "copy",
        config=value,
        token=imported["token"],
        decisions=decisions,
        digest=preview["digest"],
        start=1,
        stop=2,
    )
    assert response.json()["rows"] == [["Oct-26"], [""], ["Sep-26"], ["Oct-26"]]


def test_import_xlsx_sheet_and_header_row(admin_client: Client, placed: Placement) -> None:
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.append(["Wrong sheet"])
    sheet = workbook.create_sheet("Размещения")
    sheet.append(["Report title"])
    sheet.append([None, "Месяц", "Месяц"])
    sheet.append(["example.com", "Oct-26", "Sep-26"])
    stream = io.BytesIO()
    workbook.save(stream)
    response = admin_client.post(
        reverse("admin:placements_placement_export_designer"),
        {
            "action": "import",
            "sheet": "1",
            "header_row": "2",
            "file": SimpleUploadedFile("book.xlsx", stream.getvalue()),
        },
    )
    result = response.json()
    assert result["sheets"] == ["Sheet", "Размещения"]
    assert [c["header"] for c in result["config"]["columns"]] == ["", "Месяц", "Месяц"]
    assert result["count"] == 1


def test_profiles_are_personal_and_revision_protected(
    admin_client: Client, placed: Placement
) -> None:
    value = config("target")
    saved = post(admin_client, "save", config=value, name="Report").json()
    assert post(admin_client, "save", config=value, name="report").status_code == 409
    user = User.objects.create_superuser("other", "other@example.com", "test")
    other = Client()
    other.force_login(user)
    assert (
        post(other, "save", config=value, name="Report", id=saved["id"], revision=1).status_code
        == 403
    )
    changed = post(admin_client, "save", config=value, name="Renamed", id=saved["id"], revision=1)
    assert changed.json()["revision"] == 2
    assert (
        post(admin_client, "save", config=value, name="Old", id=saved["id"], revision=1).status_code
        == 409
    )
    deleted = post(admin_client, "delete", id=saved["id"], revision=2).json()
    assert deleted["profiles"] == []
    assert (
        post(admin_client, "restore", id=saved["id"], revision=3).json()["profiles"][0]["name"]
        == "Renamed"
    )


def test_draft_restores_for_same_user(admin_client: Client, placed: Placement) -> None:
    value = config("target")
    assert post(admin_client, "draft", config=value).status_code == 200
    assert ExportProfile.objects.get().name == ""
    page = admin_client.get(reverse("admin:placements_placement_export_designer"))
    assert page.context["bootstrap"]["config"] == value


def test_selection_and_filter_are_preserved(admin_client: Client, placed: Placement) -> None:
    other = Placement.objects.create(
        site=Site.objects.create(domain="other.com"), product=placed.product
    )
    url = reverse("admin:placements_placement_export_designer") + "?q=example.com"
    response = admin_client.post(
        url,
        {
            "action": "preview",
            "ids": [str(other.pk)],
            "payload": json.dumps({"config": config("target")}),
        },
    )
    assert response.json()["count"] == 0


def test_list_page_size_in_address(admin_client: Client, placed: Placement) -> None:
    # «На странице: 500» в «Размещениях» — параметр адреса, а не отбор (ошибка 20.10.2026).
    url = reverse("admin:placements_placement_export_designer") + "?per_page=500&o=9"
    assert admin_client.post(url, {"action": "open"}).status_code == 200
    response = admin_client.post(
        url, {"action": "preview", "payload": json.dumps({"config": config("target")})}
    )
    assert response.json()["count"] == 1


def test_unknown_fields_and_forged_template_token_rejected(
    admin_client: Client, placed: Placement
) -> None:
    assert post(admin_client, "preview", config=config("sql:select *")).status_code == 400
    assert post(admin_client, "preview", config=config(), token="wrong").status_code == 400
    response = admin_client.get(reverse("admin:placements_placement_export_designer"))
    assert response.status_code == 200


def test_constructor_keeps_admin_navigation(admin_client: Client, placed: Placement) -> None:
    response = admin_client.get(reverse("admin:placements_placement_export_designer"))
    assert response.status_code == 200
    assert response.context["is_nav_sidebar_enabled"]
    assert response.context["available_apps"]
    assert response.context["opts"] == Placement._meta
    html = response.content.decode()
    assert 'id="header"' in html and 'id="nav-sidebar"' in html
    assert 'id="palette-switch"' in html


class TestCurrency:
    """Валюта сумм: «Исходная валюта» и «Пересчитать в» (просьба 19.10.2026)."""

    @pytest.fixture(autouse=True)
    def rates(self) -> None:
        for code, rate in (("USD", "1.25"), ("GBP", "0.8")):
            ExchangeRate.objects.create(
                currency=code, rate_date=dt.date(2026, 10, 16), rate=Decimal(rate)
            )

    def _rows(self, client: Client, value: dict[str, Any]) -> list[list[str]]:
        response = post(client, "preview", config=value)
        assert response.status_code == 200, response.content
        rows: list[list[str]] = response.json()["rows"]
        return rows

    def test_dollar_payment_in_euro(self, admin_client: Client, placed: Placement) -> None:
        placed.price_paid_cents, placed.currency = 10000, "USD"
        placed.save()
        value = config("paid", "currency", "paid_eur")
        assert self._rows(admin_client, value) == [["100.00", "USD", "80.00"]]
        value["currency_to"] = "EUR"
        assert self._rows(admin_client, value) == [["80.00", "EUR", "80.00"]]

    def test_each_payment_from_its_own_currency(
        self, admin_client: Client, placed: Placement
    ) -> None:
        # «Как в размещении»: евро остаются евро, в фунты — по курсу.
        value = {**config("paid", "currency"), "currency_to": "GBP"}
        assert self._rows(admin_client, value) == [["98.76", "GBP"]]

    def test_chosen_source_overrides_recorded(
        self, admin_client: Client, placed: Placement
    ) -> None:
        # Записано в евро, но выбрано «все суммы в долларах» — 123.45 считаются долларами.
        value = {**config("paid", "currency", "paid_eur"), "currency_from": "USD"}
        assert self._rows(admin_client, value) == [["123.45", "USD", "98.76"]]

    def test_site_prices_and_default_header(self) -> None:
        registry = design.fields()
        value = {**config("placement_price", "announce_price"), "currency_to": "USD"}
        value["columns"][0]["header"] = registry["placement_price"].header
        records = [{"pk": 1, "values": {"placement_price": 10000, "announce_price": None}}]
        design.convert(records, value)
        assert records[0]["values"] == {"placement_price": 12500, "announce_price": None}
        plan = design.build_plan(records, design.validate(value, registry), registry)
        # Стандартный заголовок — в выбранной валюте, свой — как вписан.
        assert [c.title for c in plan.sheet.columns] == [
            "Цена размещения статья, USD",
            "announce_price",
        ]

    def test_old_template_without_currency(self, admin_client: Client, placed: Placement) -> None:
        value = config("paid")
        del value["currency_from"], value["currency_to"]
        assert self._rows(admin_client, value) == [["123.45"]]

    def test_unknown_currency_rejected(self, admin_client: Client, placed: Placement) -> None:
        value = {**config("paid"), "currency_to": "XXX"}
        assert post(admin_client, "preview", config=value).status_code == 400

    def test_choice_list(self, admin_client: Client, placed: Placement) -> None:
        response = admin_client.get(reverse("admin:placements_placement_export_designer"))
        assert response.context["bootstrap"]["currencies"] == ["EUR", "USD", "GBP"]


@pytest.mark.parametrize(
    ("cents", "source", "target", "expected"),
    [
        (10000, "USD", "USD", 10000),
        (10000, "USD", "EUR", 8000),
        (10000, "EUR", "GBP", 8000),
        (10000, "USD", "GBP", 6400),
        (1, "EUR", "USD", 1),
        (10000, "USD", "JPY", None),
    ],
)
def test_convert_cents(cents: int, source: str, target: str, expected: int | None) -> None:
    rates = {"EUR": Decimal(1), "USD": Decimal("1.25"), "GBP": Decimal("0.8")}
    assert convert_cents(cents, source, target, rates) == expected
